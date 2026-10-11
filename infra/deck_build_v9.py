"""Colosseum deck v9 (2026-10-11): po Nemanjinom production brief-u od 11.10.
Sedam slajdova, marketplace za pcelinje proizvode od pcelara od poverenja.
Naslovi i podnaslovi su doslovno iz brief-a. Stil kao v6-v8 (YC: jedna ideja
po slajdu, recenica u naslovu; redosled po Pixar strukturi), tamna pozadina,
zuti naslovi. Mesta za prave fotografije su isprekidani okviri sa opisom;
Nemanja ih zamenjuje u Slides-u (mreza u cloud sesiji ne dozvoljava
preuzimanje sa Drive-a, futard.io ni hivebits.io).
Tekst je i u content/colosseum/deck.md (izvor istine za tekst i izvore).
Pokretanje: python3 infra/deck_build_v9.py <izlaz.pptx>
Ikonice: Noto Emoji v2.042 (Google, Apache 2.0), u infra/deck_icons/."""
import sys, os
ICONS=os.path.join(os.path.dirname(os.path.abspath(__file__)),'deck_icons')
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
from lxml import etree

Y=RGBColor(0xEF,0xB7,0x3E); W=RGBColor(0xFF,0xFF,0xFF); G=RGBColor(0x9A,0x9A,0xA0)
L=RGBColor(0x8A,0x8A,0x90); D=RGBColor(0x2A,0x2A,0x2E)
FONT='Poppins'
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
C,Lf,Rt=PP_ALIGN.CENTER,PP_ALIGN.LEFT,PP_ALIGN.RIGHT

def slide(notes=''):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    f=s.background.fill; f.gradient(); f.gradient_angle=315
    f.gradient_stops[0].color.rgb=RGBColor(0x1C,0x1C,0x1F); f.gradient_stops[1].color.rgb=RGBColor(0x38,0x37,0x3C)
    if notes: s.notes_slide.notes_text_frame.text=notes
    return s

def tb(s,x,y,w,h,runs,size=20,color=W,bold=False,align=C,anchor=MSO_ANCHOR.TOP,font=FONT):
    t=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)).text_frame
    t.word_wrap=True; t.vertical_anchor=anchor
    if isinstance(runs,str): runs=[(runs,color)]
    p=t.paragraphs[0]; p.alignment=align
    for text,col in runs:
        parts=text.split('\n')
        for i,part in enumerate(parts):
            if i>0: p=t.add_paragraph(); p.alignment=align
            if not part: continue
            r=p.add_run(); r.text=part; r.font.size=Pt(size); r.font.bold=bold
            r.font.color.rgb=col
            if font: r.font.name=font
    return t

def icon_pic(s,name,cx,y,size=1.25):
    s.shapes.add_picture(os.path.join(ICONS,name+'.png'),Inches(cx-size/2),Inches(y),Inches(size),Inches(size))

def arrow(s,x1,y,x2,color=L):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y),Inches(x2),Inches(y))
    c.line.color.rgb=color; c.line.width=Pt(1.5)
    ln=c.line._get_or_add_ln(); ln.append(etree.SubElement(ln,qn('a:tailEnd'),type='triangle'))

def photo(s,x,y,w,h,label):
    """Mesto za pravu fotografiju: isprekidan okvir sa opisom. Zamenjuje se u Slides-u."""
    r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb=D
    r.line.color.rgb=L; r.line.width=Pt(1.25); r.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    r.shadow.inherit=False
    tb(s,x+0.2,y+h/2-0.5,w-0.4,1.0,label,15,G,False,C,MSO_ANCHOR.MIDDLE)

def box(s,x,y,w,h,title,body,title_color=Y):
    r=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb=D; r.line.color.rgb=L; r.line.width=Pt(1); r.shadow.inherit=False
    tb(s,x+0.3,y+0.25,w-0.6,0.6,title,22,title_color,True,Lf)
    tb(s,x+0.3,y+0.95,w-0.6,h-1.1,body,20,W,False,Lf)

def head1(s,text,size=36,w=11.7): tb(s,0.8,0.55,w,1.7,text,size,Y,True,Lf)
def sub(s,text,y=2.25,size=22,w=11.7): tb(s,0.8,y,w,0.8,text,size,W,False,Lf)
def big(s,y,text,size=28,color=W,bold=True,x=0.8,w=11.7,h=0.9): tb(s,x,y,w,h,text,size,color,bold,Lf)
def foot(s,text,color=G): tb(s,0.8,6.8,11.7,0.4,text,13,color,False,Lf)

# 1 cover
s=slide('Hi, I am Nemanja. We are HiveBits. HiveBits connects buyers directly with bee products from beekeepers they can trust. '
        'We start with beekeepers we know. Later, our hardware brings more beekeepers into the network.\n\n'
        'A: HiveBits presented smart hives and farm ownership at Colosseum before (colosseum.com/projects/explore/hivebits, /hivebits-1). '
        'B: existing beekeeping operation, smart-hive technology, marketplace business model (brief 11.10; marketplace treated as existing product for this draft). '
        'C: extending into direct commerce between buyers and trusted beekeepers.\n\n'
        'FOTO: desna polovina je mesto za pravu fotografiju pcelinjaka ili rada sa kosnicama (Drive hivebits/Swarm). Ne AI slika.')
tb(s,0.8,0.7,6.4,1.0,'HiveBits',44,W,True,Lf)
tb(s,0.8,1.9,6.4,2.6,'HiveBits connects buyers directly with bee products from beekeepers they can trust.',32,Y,True,Lf)
tb(s,0.8,4.7,6.4,1.4,'Start with beekeepers we know.\nScale the network through HiveBits hardware.',20,W,False,Lf)
photo(s,7.6,0.7,5.0,6.1,'Photo: HiveBits beekeeping operation\n(real photo from Drive)')

# 2 problem
s=slide('[Problem]\n\nWhen you buy honey, you usually see a label. You do not see the beekeeper. We think trust starts with the person who makes the honey. '
        'So we put the beekeeper and where the honey comes from in front, not just the product.\n\n'
        'A: direct-to-consumer agricultural commerce already exists (drenio.rs, crowdfarming.com); we do not claim to be the only one. '
        'B: HiveBits starts from practical beekeeping experience and its own farm. C: make the beekeeper and product-origin information part of the buying experience.\n\n'
        'NE GOVORI: "nobody knows where their food comes from", "all intermediaries are unnecessary" (brief).')
head1(s,'Buyers should know the beekeeper behind the jar.')
sub(s,'Trust begins with the people who produce the honey.')
YI=3.4
icon_pic(s,'farmer',3.0,YI,1.5); tb(s,1.6,YI+1.6,2.8,0.5,'The beekeeper',20,W,True)
jar=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(6.05),Inches(YI+0.25),Inches(1.2),Inches(1.25))
jar.fill.solid(); jar.fill.fore_color.rgb=Y; jar.line.fill.background(); jar.shadow.inherit=False
lid=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(6.15),Inches(YI),Inches(1.0),Inches(0.3))
lid.fill.solid(); lid.fill.fore_color.rgb=W; lid.line.fill.background(); lid.shadow.inherit=False
tb(s,5.25,YI+1.6,2.8,0.5,'The jar',20,W,True)
icon_pic(s,'customer',10.3,YI,1.5); tb(s,8.9,YI+1.6,2.8,0.5,'The buyer',20,W,True)
arrow(s,3.9,YI+0.75,5.9,Y); arrow(s,7.4,YI+0.75,9.4,Y)

# 3 solution
s=slide('[Solution]\n\nThe marketplace is simple. You find a beekeeper. You see what he has. You buy directly from him. That is it. No extra steps.\n\n'
        'A: first smart-hive product already sold (futard.io launch page). B: for this draft the marketplace is treated as an existing product (founder instruction). '
        'CONFIRM THE REAL INTERFACE BEFORE SUBMISSION (open-questions #47). C: demonstrate a complete buying journey (demo video).\n\n'
        'NE PISE: escrow, delivery tracking, lab certification, smart-contract settlement, automated payments, osim ako stvarno postoje (brief).\n'
        'SCREENSHOT: desna polovina je mesto za screenshot pravog marketplace interfejsa.')
head1(s,'Discover a beekeeper. Choose a product. Buy directly.',30,6.6)
sub(s,'A straightforward marketplace for bee products.',2.5,20,6.6)
steps=[('1','Discover a beekeeper.'),('2','Browse available bee products.'),('3','Buy directly.')]
for i,(n,t) in enumerate(steps):
    y=3.4+i*1.05
    c=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(0.85),Inches(y),Inches(0.7),Inches(0.7))
    c.fill.solid(); c.fill.fore_color.rgb=Y; c.line.fill.background(); c.shadow.inherit=False
    tb(s,0.85,y+0.08,0.7,0.6,n,22,D,True)
    tb(s,1.75,y+0.08,5.5,0.7,t,24,W,True,Lf)
photo(s,7.6,0.7,5.0,6.1,'Screenshot: HiveBits marketplace\n(real interface, confirm before submission)')

# 4 network
s=slide('[How the network grows]\n\nWe do not need every beekeeper on day one. We start with people we know and work with. Our hardware gives us a way to bring in new beekeepers later, '
        'with monitoring data from their hives. The hardware does not replace trust. It helps us keep the connection to real production.\n\n'
        'A: own beekeeping operation and smart-hive technology exist (post @hivebits_io 05.09.2026, "every colony monitored by our own hardware"). '
        'B: the initial network is beekeepers known to and working with the team. C: hardware as a path for other beekeepers to join, no date.\n\n'
        'OGRADA (brief, pravilo 9): hardware gives monitoring data; it does not prove honey quality, authenticity or beekeeper reliability. Nema "verified", "certified", "guaranteed".')
head1(s,'We start with beekeepers we know. Hardware opens the network over time.',32)
sub(s,'Start with trusted relationships. Expand through a repeatable onboarding model.',2.3,20)
box(s,0.8,3.4,5.6,2.9,'Stage 1','HiveBits and the beekeepers we know.\nOur own farm is the first supplier.')
box(s,7.0,3.4,5.6,2.9,'Stage 2','More beekeepers join the network using HiveBits hardware.\nMonitoring data keeps the link to real production.')
arrow(s,6.5,4.85,6.9,Y)
foot(s,'Hardware provides monitoring data. Trust still comes from the people.')

# 5 traction
s=slide('[Proof of demand]\n\nBefore the marketplace we sold our first product, the smart hive. 150 people backed 300 smart hives in 27 hours. '
        'That money builds the farm now, and we record every step. This farm and these beekeepers are the base for the marketplace.\n\n'
        'A: 150 people, 300 smart hives, 27 hours. Izvori: screenshot MetaDAO stranice (news.md 10.10.2026: 150 funders, $140K cilj, $243K ponudjeno, 173%); '
        'post @hivebits_io 08.09.2026 ("one day and three hours"); futard.io/launch/AbTuTsoBJS2jQggvfb7d1LA2SfbHa5NB5HbJqX8C2PnS. '
        'B: funds build the real farm, progress documented (klip 10, postovi od 16.09). C: existing operation as the base for the marketplace.\n\n'
        'Brojke tacno kako su potvrdjene. Bez prihoda, konverzije, retencije, profitabilnosti. Dolarski iznosi nisu na slajdu (brief ih ne trazi, #49).\n'
        'SLIKE: levo screenshot Futard launch stranice, desno prava fotka farme ili kosnica.')
head1(s,'150 supporters. 300 smart hives. 27 hours.',40)
sub(s,'Our first product demonstrated demand for real-world beekeeping infrastructure.',2.3,20)
photo(s,0.8,3.3,5.6,3.3,'Screenshot: Futard launch page\n(futard.io/launch/..., real)')
photo(s,7.0,3.3,5.6,3.3,'Photo: the farm or smart hives\n(real photo from Drive)')

# 6 markets
s=slide('[Markets]\n\nWe start with honey. We look at four markets where we have beekeepers we know: Brazil, Texas, Montreal and the European Union. '
        'We will test demand, supply and logistics there. We take 5 percent of each sale. The beekeeper chooses if he pays it or adds it to his price.\n\n'
        'A: initial demand for the smart-hive product shown (slide 5). B: the team identified Brazil, Texas, Montreal and the EU as target markets (brief; Nemanja 24.09). TARGET, not existing. '
        'C: test demand, supply availability and logistics across the four markets.\n\n'
        '5%: Nemanja 24.09 and brief. TAM: NO NUMBER ON THE SLIDE until a sourced bottom-up is confirmed (brief). Only sourced input available: US, NielsenIQ for National Honey Board, '
        '52 weeks to 2 Nov 2024, 46.2M households x $18.15 = $839M a year (honey.com/images/files/National-Honey-Board-Category-Review-Analysis-2024.pdf). Whole US, not Texas; not on the slide until Nemanja says so (#50). '
        'Brazil, Montreal, EU: no equivalent yet (researcher R4).')
head1(s,'Honey is our starting category. Four markets are in our sights.',34)
for i,m in enumerate(['Brazil','Texas','Montreal','European Union']):
    x=0.8+i*2.95
    r=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(2.6),Inches(2.75),Inches(1.5))
    r.fill.solid(); r.fill.fore_color.rgb=D; r.line.color.rgb=Y; r.line.width=Pt(1.25); r.shadow.inherit=False
    tb(s,x,2.6,2.75,1.5,m,24,W,True,C,MSO_ANCHOR.MIDDLE)
tb(s,0.8,4.25,11.7,0.5,'Target markets. We test demand, supply and logistics there.',18,G,False,Lf)
big(s,5.1,'Marketplace fee: 5% of each sale.',30,Y)
big(s,5.85,'The beekeeper chooses to pay it or add it to the price.',22,W,False)

# 7 next milestone and team
s=slide('[Next milestone]\n\nOur next milestone is repeat sales. Not one sale, but people who come back to the same beekeeper. For that we build the marketplace, onboard the first beekeepers, '
        'connect each product to where it comes from, and test sales in the four markets. I am a beekeeper and an engineer, and I run an apitherapy centre. Sinisa manages the farm. Thank you.\n\n'
        'A: existing operation, smart-hive technology, initial demand (slide 5). B: 5% marketplace fee (slide 6). C: the four priorities below are development priorities, NOT finished today (brief).\n\n'
        'Tim: samo potvrdjene cinjenice iz brief-a. Bez ask-a, iznosa i obecanja prinosa (ask nije potvrdjen, #51).\n'
        'FOTO: desno dole mesto za pravu fotografiju osnivaca.')
head1(s,'Our next milestone is repeatable direct sales from trusted beekeepers.',32)
sub(s,'Build the marketplace. Onboard the first beekeeper network. Validate repeat purchases.',2.3,19)
tb(s,0.8,3.2,7.0,0.5,'Development priorities',18,Y,True,Lf)
tb(s,0.8,3.65,7.0,2.8,'1.  Marketplace and the full buyer journey.\n2.  First beekeepers and product listings.\n3.  Product-origin information linked to available farm data.\n4.  Demand and direct sales tests in the target markets.',18,W,False,Lf)
tb(s,8.2,3.2,4.4,0.5,'Team',18,Y,True,Lf)
tb(s,8.2,3.6,4.4,2.0,[('Nemanja Šćepanović\n',W),('Beekeeper and engineer. Runs an apitherapy centre.\n\n',G),('Siniša Kežić\n',W),('Manages the farm.',G)],15,W,False,Lf)
photo(s,8.2,5.7,4.4,1.1,'Photo: founders (real)')

prs.save(sys.argv[1]); print(len(prs.slides._sldIdLst),'slides')
