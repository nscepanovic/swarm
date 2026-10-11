"""Colosseum deck v10 (2026-10-11): po Nemanjinom "Complete Pitch Deck Creation
Brief" od 11.10. Osam slajdova, tri stuba (participate, reserve, buy),
konkurencija kao matrica, bottom-up trziste po cetiri trzista.
Stil kao v6-v9 (YC: jedna ideja po slajdu, recenica u naslovu; Pixar redosled),
tamna pozadina, zuti naslovi. Mesta za prave fotografije su isprekidani okviri.
Tekst je i u content/colosseum/deck.md (izvor istine za tekst i izvore).
Pokretanje: python3 infra/deck_build_v10.py <izlaz.pptx>
Ikonice: Noto Emoji v2.042 (Google, Apache 2.0), u infra/deck_icons/."""
import sys, os
ICONS=os.path.join(os.path.dirname(os.path.abspath(__file__)),'deck_icons')
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
from lxml import etree

Y=RGBColor(0xEF,0xB7,0x3E); W=RGBColor(0xFF,0xFF,0xFF); G=RGBColor(0x9A,0x9A,0xA0)
L=RGBColor(0x8A,0x8A,0x90); D=RGBColor(0x2A,0x2A,0x2E); GR=RGBColor(0x6F,0xC2,0x8A); BL=RGBColor(0x7F,0xB3,0xE6)
FONT='Poppins'
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
C,Lf,Rt=PP_ALIGN.CENTER,PP_ALIGN.LEFT,PP_ALIGN.RIGHT

def slide(notes=''):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    f=s.background.fill; f.gradient(); f.gradient_angle=315
    f.gradient_stops[0].color.rgb=RGBColor(0x1C,0x1C,0x1F); f.gradient_stops[1].color.rgb=RGBColor(0x38,0x37,0x3C)
    if notes: s.notes_slide.notes_text_frame.text=notes
    return s

def tb(s,x,y,w,h,runs,size=20,color=W,bold=False,align=C,anchor=MSO_ANCHOR.TOP,font=FONT):
    t=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)).text_frame
    t.word_wrap=True; t.vertical_anchor=anchor
    if isinstance(runs,str): runs=[(runs,color)]
    p=t.paragraphs[0]; p.alignment=align
    for text,col in runs:
        parts=text.split('\n')
        for i,part in enumerate(parts):
            if i>0: p=t.add_paragraph(); p.alignment=align
            if not part: continue
            r=p.add_run(); r.text=part; r.font.size=Pt(size); r.font.bold=bold
            r.font.color.rgb=col
            if font: r.font.name=font
    return t

def icon_pic(s,name,cx,y,size=1.25):
    s.shapes.add_picture(os.path.join(ICONS,name+'.png'),Inches(cx-size/2),Inches(y),Inches(size),Inches(size))

def arrow(s,x1,y,x2,color=L):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y),Inches(x2),Inches(y))
    c.line.color.rgb=color; c.line.width=Pt(1.5)
    ln=c.line._get_or_add_ln(); ln.append(etree.SubElement(ln,qn('a:tailEnd'),type='triangle'))

def rect(s,x,y,w,h,fill=D,line=L,lw=1.0,rounded=True,dash=False):
    r=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    if fill is None: r.fill.background()
    else: r.fill.solid(); r.fill.fore_color.rgb=fill
    if line is None: r.line.fill.background()
    else:
        r.line.color.rgb=line; r.line.width=Pt(lw)
        if dash: r.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    r.shadow.inherit=False
    return r

def photo(s,x,y,w,h,label):
    """Mesto za pravu fotografiju: isprekidan okvir sa opisom. Zamenjuje se u Slides-u."""
    rect(s,x,y,w,h,D,L,1.25,False,True)
    tb(s,x+0.2,y+h/2-0.5,w-0.4,1.0,label,14,G,False,C,MSO_ANCHOR.MIDDLE)

def head1(s,text,size=34,w=11.7): tb(s,0.8,0.5,w,1.6,text,size,Y,True,Lf)
def sub(s,text,y=2.2,size=20,w=11.7,color=W): tb(s,0.8,y,w,0.9,text,size,color,False,Lf)
def foot(s,text,color=G,size=13): tb(s,0.8,6.85,11.7,0.4,text,size,color,False,Lf)
def tag(s,x,y,text,color):
    r=rect(s,x,y,1.4,0.34,None,color,1.0)
    tb(s,x,y+0.02,1.4,0.3,text,10,color,True,C)

# 1 cover
s=slide('Hi, I am Nemanja. We are HiveBits. HiveBits connects people directly with real farms. You can take part in production, reserve what a farm will make, '
        'or buy directly from a farmer you can trust. We started with bees. The model is built for other farms too.\n\n'
        'A: two earlier Colosseum entries pitched the smart hive (brief). B: beekeeping operation, farm under construction in Bosnia, hive sensors (brief). '
        'C: reserve and buy as product, other categories (brief: do not imply all categories are operational).\n\n'
        'FOTO: desno je mesto za pravu fotografiju HiveBits kosnica (Drive hivebits/Swarm). Ne AI slika.')
tb(s,0.8,0.6,6.6,0.9,'HiveBits',40,W,True,Lf)
tb(s,0.8,1.5,6.6,1.6,'HiveBits connects people directly with real farms.',32,Y,True,Lf)
tb(s,0.8,3.2,6.6,1.5,'Participate in production. Reserve future output. Buy directly from farmers you can trust.',18,W,False,Lf)
tb(s,0.8,4.5,6.6,0.6,'First: beekeeping. Next: other farms that produce.',16,G,False,Lf)
# diagram people -> farms -> products -> people
yd=5.5
for cx,icon,label in [(1.6,'customer','People'),(4.0,'farmer','Farms'),(6.4,'seedling','Products')]:
    icon_pic(s,icon,cx,yd,0.8); tb(s,cx-1.0,yd+0.85,2.0,0.4,label,13,W,True)
arrow(s,2.1,yd+0.4,3.5,Y); arrow(s,4.5,yd+0.4,5.9,Y)
tb(s,0.8,yd-0.45,6.6,0.4,'HiveBits',12,Y,True,Lf)
photo(s,7.7,0.6,4.9,6.3,'Photo: real HiveBits hives\n(from Drive)')

# 2 problem
s=slide('[Problem]\n\nBetween a farm and a customer there are usually a few companies. The farmer sees only the first one, and that one sets his price. '
        'In the US, a beekeeper gets $2.45 a pound from a packer. If he sells it himself, $7.15. Same honey. People who fund farms, people who buy, and the farmer are in three different worlds. We want to put them in one place.\n\n'
        'A: USDA NASS Honey, March 2026: 2025 average price sold to co-ops/packers $2.45/lb; retail (beekeeper sells direct) $7.15/lb. https://esmis.nal.usda.gov/sites/default/release-files/795818/hony0326.txt (market.md 4a). '
        'One category, one country, one primary source, as the brief requires. C: "shortens the distance" is intent.\n\nNE GOVORI: "all intermediaries are unnecessary", "every farmer gets one third".')
head1(s,'Agriculture needs a better connection between production and customers.',32)
sub(s,'Production, financing, distribution and retail are separate worlds. The farmer meets only the first one.',2.2,19)
yc=3.5
nodes=[(1.7,'farmer','Beekeeper'),(5.0,'factory','Packer'),(8.3,'store','Store'),(11.6,'customer','Customer')]
for x,icon,label in nodes:
    icon_pic(s,icon,x,yc,1.1); tb(s,x-1.3,yc+1.15,2.6,0.4,label,16,W,True)
for i in range(3): arrow(s,nodes[i][0]+0.75,yc+0.55,nodes[i+1][0]-0.75)
tb(s,0.4,yc+1.6,2.6,0.6,'$2.45 / lb',26,Y,True)
tb(s,10.3,yc+1.6,2.6,0.6,'$7.15 / lb',26,Y,True)
tb(s,3.0,yc+1.7,7.3,0.5,'US honey, 2025. Same beekeeper: to a packer, or sold direct.',15,G,False)
tb(s,0.8,6.1,11.7,0.6,'HiveBits shortens the distance between the farm and the people who eat what it makes.',20,W,True,Lf)
foot(s,'Source: USDA NASS Honey report, March 2026, 2025 average prices.')

# 3 three pillars
s=slide('[Product: three pillars]\n\nThree ways. One, you help fund real production. We did this once, with our own farm: 300 smart hives. Two, you rent a hive or a cow for a season, '
        'or reserve part of a future harvest. The farmer plans better and you know where your food comes from. This part is planned. Three, you buy honey and other products directly from farmers we know. '
        'We are building this now. We take 5 percent of each sale. These are three different deals with different terms. We do not mix them.\n\n'
        'A (PARTICIPATE): 300 hives, 150 people, 27 hours, $140K target, Futard (brief section 4). "Help fund", not "own" or "invest": ownership rights and legal structure not confirmed (#43). '
        'C (RESERVE): brief: treat as product direction. Renting a unit, reserving future output and investing are separate, as the brief requires. '
        'B/C (BUY): being built; first beekeepers are people the team knows (brief). 5%: brief and Nemanja 24.09.\n\nNE PISE: returns, ownership, guarantees, escrow, smart-contract settlement.')
head1(s,'One platform. Three ways to participate.',36)
cols=[('PARTICIPATE','DONE ONCE',GR,'Help fund real production.\n\nFirst: 300 smart hives, our farm.\nNext: other farms.'),
      ('RESERVE','PLANNED',BL,'Rent a hive or a cow for a season, or reserve part of a future harvest.\n\nThe farmer plans ahead. You know your source.'),
      ('BUY','BUILDING',Y,'Buy honey and other products directly from farmers we know.\n\n5% fee per sale. The farmer absorbs it or adds it to the price.')]
for i,(t,st,col,body) in enumerate(cols):
    x=0.8+i*3.95
    rect(s,x,2.1,3.75,3.9,D,L,1.0)
    tb(s,x+0.25,2.25,3.3,0.5,t,22,W,True,Lf)
    tag(s,x+0.3,2.8,st,col)
    tb(s,x+0.3,3.3,3.15,2.6,body,16,W,False,Lf)
tb(s,0.8,6.2,11.7,0.6,'Three different deals: funding, renting or reserving, and buying each come with their own terms.',16,G,False,Lf)

# 4 trust and network
s=slide('[Trust and growth]\n\nTrust comes from people first. We start with beekeepers we know and work with. Our farm is the first supplier. We already have hardware under our hives. It measures weight, temperature and humidity. '
        'Later it gives other beekeepers a way to join and show their production. To be clear: the sensor shows what happens in the hive. It does not prove the honey is good. The beekeeper does.\n\n'
        'A/B: sensors under the hive measure weight, temperature, humidity (brief; post @hivebits_io "every colony monitored by our own hardware"). B: initial network is beekeepers known to the team (brief). '
        'C: hardware as onboarding path for others (brief: intended).\n\nFOTO: desno mesto za pravu fotografiju pametne kosnice sa senzorom.')
head1(s,'Start with beekeepers we know. Scale through HiveBits hardware.',32,7.0)
rect(s,0.8,2.3,6.6,1.6,D,L,1.0)
tb(s,1.1,2.4,6.0,0.5,'Stage 1',18,Y,True,Lf)
tb(s,1.1,2.9,6.0,1.0,'Beekeepers we know and work with. Our farm is the first supplier.',16,W,False,Lf)
rect(s,0.8,4.1,6.6,1.9,D,L,1.0)
tb(s,1.1,4.2,6.0,0.5,'Stage 2',18,Y,True,Lf)
tb(s,1.1,4.7,6.0,1.3,'Other beekeepers join with HiveBits hardware under their hives: weight, temperature, humidity.',16,W,False,Lf)
tb(s,0.8,6.15,6.6,0.7,'Sensor data supports monitoring. It does not certify honey quality, authenticity or a farmer.',14,G,False,Lf)
photo(s,7.9,2.3,4.7,4.4,'Photo: smart hive with sensor\n(real)')

# 5 competition matrix
s=slide('[Competition]\n\nWe are not the first to connect buyers and farmers. Pollenity does smart hives and hive adoption. CrowdFarming lets you adopt a tree and get the harvest. Drenio sells local food direct. '
        'AgriDex trades commodities between businesses on Solana. AcreTrader and FarmTogether sell farmland investments. All of them work. What we are building is the combination: fund production, reserve it, buy it, and see the farm through hardware. '
        'We did not find that in one place. That is a hypothesis we have to prove, not a claim.\n\n'
        'IZVORI (web pretraga 11.10, sajtovi blokirani iz cloud sesije): Pollenity: EBRD news 2026 (Adopt-a-Hive subscription, single-origin honey, each batch analysed), CORDIS 887791 (BeeBot, HiHive). '
        'CrowdFarming: Euronews 05.06.2026 (adoptions, 300+ producers), FreshPlaza; fee 16-32% (market.md). Drenio: NOT VERIFIED (site blocked, no search results); description is the category from the brief (#52). '
        'AgriDex: The Defiant ($9M stablecoin trades, B2B, Africa), The Block (first RWA trade on Solana, 0.15% per side). AcreTrader: accredited only, $10-25K minimum (2026 review). FarmTogether: accredited, $15K crowdfunding minimum, $50K fund (farmtogether.com). '
        'Competitor metrics stay in notes, not on the slide. "As far as we found": competition.md 24.09 plus this search; Frontier 2026 corpus not searched.')
head1(s,'Existing platforms validate the market. HiveBits aims to connect the farm lifecycle.',28)
colsx=[4.6,6.2,7.8,9.4,11.0]; heads=['Fund\nproduction','Reserve\nor rent','Buy\ndirect','Hardware\nmonitoring','Beyond one\ncategory']
for x,h in zip(colsx,heads): tb(s,x-0.75,1.85,1.5,0.7,h,12,G,True,C)
rows=[('Pollenity','smart hives, hive adoption, honey',['-','y','y','y','-']),
      ('CrowdFarming','adopt a tree or animal, direct boxes',['-','y','y','-','y']),
      ('Drenio','local farmers to consumers',['-','-','y','-','y']),
      ('AgriDex','B2B commodity trade on Solana',['-','-','b','-','y']),
      ('AcreTrader','farmland investment, accredited',['y','-','-','-','-']),
      ('FarmTogether','farmland investment, accredited',['y','-','-','-','-']),
      ('HiveBits','beekeeping first, then other farms',['d','p','p','y','p'])]
y0=2.6; rh=0.52
for i,(name,desc,cells) in enumerate(rows):
    y=y0+i*rh
    if name=='HiveBits': rect(s,0.7,y-0.04,11.3,rh,RGBColor(0x3A,0x36,0x28),Y,1.0)
    tb(s,0.8,y,1.9,rh,name,14,Y if name=='HiveBits' else W,True,Lf,MSO_ANCHOR.MIDDLE)
    tb(s,2.5,y,2.1,rh,desc,10,G,False,Lf,MSO_ANCHOR.MIDDLE)
    for x,c in zip(colsx,cells):
        cy=y+rh/2-0.14
        if c=='y': rect(s,x-0.14,cy,0.28,0.28,W,None)
        elif c=='d': rect(s,x-0.14,cy,0.28,0.28,GR,None)
        elif c=='p': rect(s,x-0.14,cy,0.28,0.28,None,Y,1.5)
        elif c=='b': rect(s,x-0.14,cy,0.28,0.28,W,None); tb(s,x-0.6,cy+0.26,1.2,0.25,'B2B only',9,G,False,C)
        else: tb(s,x-0.3,cy-0.06,0.6,0.3,'-',14,L,False,C)
# legend
ly=6.3
rect(s,0.8,ly+0.05,0.22,0.22,W,None); tb(s,1.1,ly,2.2,0.35,'offers today',11,G,False,Lf)
rect(s,3.0,ly+0.05,0.22,0.22,GR,None); tb(s,3.3,ly,2.4,0.35,'HiveBits, done once',11,G,False,Lf)
rect(s,5.6,ly+0.05,0.22,0.22,None,Y,1.5); tb(s,5.9,ly,3.0,0.35,'HiveBits, planned or building',11,G,False,Lf)
foot(s,'All of these work. We did not find the full combination in one place. That is the hypothesis we are building, not a claim to be first.',W,13)

# 6 traction
s=slide('[Traction]\n\nBefore the marketplace, we sold our first product: the smart hive. We asked for $140,000. 150 people backed 300 hives in 27 hours, on Futard. That money builds a real farm in Bosnia right now, and we film every step. '
        'This is proof people want to take part in real production. It is the base the marketplace starts from. It is not the marketplace\'s traction. That comes next.\n\n'
        'A: 150, 300, 27h, $140K, Futard, MetaDAO, $NECTAR: brief section 4; MetaDAO screenshot (news.md 10.10: 150 funders, 173%); post @hivebits_io 08.09.2026 ("one day and three hours"); futard.io/launch/AbTuTsoBJS2jQggvfb7d1LA2SfbHa5NB5HbJqX8C2PnS. '
        'B: farm in Bosnia under construction, documented (brief). NE PISE: $243K (#49), revenue, conversion, "marketplace traction".\n\nSLIKE: levo screenshot Futard stranice (Nemanja ga daje), desno prava fotka farme.')
head1(s,'150 supporters backed 300 smart hives in 27 hours.',36)
tb(s,0.8,2.05,11.7,0.5,'Target: $140,000. Launched on Futard, in the MetaDAO ecosystem. Token: $NECTAR.',17,W,False,Lf)
tb(s,0.8,2.55,11.7,0.5,'The money builds a real farm in Bosnia. We record every step.',17,W,False,Lf)
photo(s,0.8,3.3,5.6,3.0,'Screenshot: Futard launch page\n(provided by the founder)')
photo(s,7.0,3.3,5.6,3.0,'Photo: the farm in Bosnia\n(real)')
foot(s,'This is demand for our first product. The marketplace starts from here.',Y,15)

# 7 market
s=slide('[Market and expansion]\n\nWe start with honey because that is our farm and our people. We look at four markets where we know beekeepers: Texas, the European Union, Montreal and Brazil. '
        'The numbers on the slide are category demand, bottom-up, from official data. Texas alone is about $71M of honey a year at retail. We have not set how much of that we can reach. We will, from the first sales. '
        'Then the same model goes to other farms: cattle, milk, other things that produce. A producer, production data, people who take part, and a direct sale. That part is the vision, not a product yet.\n\n'
        'METHOD: households (or people) x yearly honey consumption x price = category demand per market. SAM and SOM not estimated (brief: do not invent; #53).\n'
        'Texas: 10,992,816 households 2020-24 (US Census QuickFacts). 35.7% of households buy honey, $18.15 a year per buyer: NielsenIQ panel for National Honey Board, 52 wks to 2 Nov 2024, whole US; ASSUMPTION: national share applied to Texas. 10.99M x 0.357 x 18.15 = $71.2M.\n'
        'EU: 436.5K t apparent consumption, 0.99 kg per person (2022, CBI); retail EUR 5.58-24.94/kg (Stiftung Warentest 4/2025, Germany), EUR 6-12 taken as typical range. 202M households (Eurostat 2024) for context.\n'
        'Montreal: 4,291,732 people CMA 2021 (StatCan via citypopulation.de, secondary, #54); 0.99 kg per person (StatCan 32-10-0054-01, 2025) = 4,249 t. Canadian retail honey price from a primary source NOT FOUND.\n'
        'Brazil: 203,062,512 people (IBGE 2022); about 60 g per person (Sebrae/A.B.E.L.H.A., ABEMEL, secondary) = 12,184 t. Brazil is a supply market: 7th producer 67K t (FAOSTAT 2024), exports $116.5M 2025; US tariff 50% and anti-dumping on raw honey (market.md 8).\n'
        'Livestock, milk: no numbers, category only (brief: may follow).')
head1(s,'Start with honey. Expand into other productive agricultural categories.',30)
tb(s,0.8,1.95,11.7,0.5,'Honey and bee products now. Then livestock, milk and dairy, other productive assets. Same model: producer, production data, participation, direct sale.',14,G,False,Lf)
mk=[('Texas','11.0M households x 35.7% buy honey\nx $18.15 a year',"about $71M a year"),
    ('European Union','436K t eaten a year\nx EUR 6 to 12 per kg retail','EUR 2.6 to 5.2B a year'),
    ('Montreal','4.3M people x 0.99 kg a year\n= about 4,200 t','retail price: data missing'),
    ('Brazil','203M people x 0.06 kg a year\n= about 12,000 t','supply market first: 7th producer')]
for i,(name,calc,res) in enumerate(mk):
    x=0.8+i*2.95
    rect(s,x,2.7,2.75,3.3,D,Y,1.25)
    tb(s,x+0.2,2.8,2.35,0.5,name,19,W,True,Lf)
    tb(s,x+0.2,3.35,2.35,1.4,calc,12,G,False,Lf)
    tb(s,x+0.2,4.9,2.35,1.0,res,15,Y,True,Lf)
tb(s,0.8,6.1,11.7,0.4,'Honey category demand per market, bottom-up. Serviceable and reachable share: not estimated yet. Set after first sales data.',13,W,False,Lf)
foot(s,'Target markets. We do not operate there yet. Sources: US Census, NielsenIQ for National Honey Board, CBI, Stiftung Warentest, Statistics Canada, IBGE, Sebrae. Details in notes.',G,11)

# 8 next phase and team
s=slide('[Next phase]\n\nWhat we build next. The main user journeys. The first network of beekeepers we know. Every listing linked to the farm and where the product comes from. Sales tests in the four markets. '
        'And a way to onboard the next producer that we can repeat. We earn 5 percent on each sale. Repeat purchases make that grow. Our milestone is simple: people buying again from the same beekeeper, in more than one market. '
        'I am a beekeeper and an engineer and I run an apitherapy centre. Sinisa manages the farm. We were at Colosseum twice with the smart hive. This time we bring the platform. Thank you.\n\n'
        'C: five priorities (brief section 8). No amounts, no allocation (ask not confirmed, #51). B: 5% (brief); other streams not priced (brief section 7). A: two earlier entries, both about the smart hive (brief). GitHub status not verified (#22).\n\nFOTO: desno dole, osnivaci.')
head1(s,'Build the marketplace that connects participation, production and commerce.',28)
tb(s,0.8,2.0,7.2,0.4,'Next',16,Y,True,Lf)
tb(s,0.8,2.4,7.2,2.6,'1.  Build and validate the main user journeys.\n2.  Onboard the first network of trusted beekeepers.\n3.  Link each listing to farm and origin data.\n4.  Test sales and demand in the four markets.\n5.  Make onboarding repeatable for the next producers.',15,W,False,Lf)
tb(s,0.8,5.0,7.2,0.9,'Revenue today comes from one confirmed rule: 5% on marketplace sales. Repeat purchases make it grow.',14,G,False,Lf)
tb(s,0.8,5.85,7.2,0.9,'Milestone: repeat purchases from the same beekeepers, in more than one market.',17,Y,True,Lf)
tb(s,8.4,2.0,4.2,0.4,'Team',16,Y,True,Lf)
tb(s,8.4,2.4,4.2,1.7,[('Nemanja Šćepanović\n',W),('Beekeeper and engineer. Runs an apitherapy centre.\n\n',G),('Siniša Kežić\n',W),('Manages the farm.',G)],14,W,False,Lf)
tb(s,8.4,4.2,4.2,1.0,'Colosseum before: Cypherpunk 2025, Frontier 2026. Both pitched the smart hive. This one is the platform.',12,G,False,Lf)
photo(s,8.4,5.4,4.2,1.4,'Photo: founders (real)')

prs.save(sys.argv[1]); print(len(prs.slides._sldIdLst),'slides')
