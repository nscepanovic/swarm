"""Colosseum deck v11 (2026-10-12): po Nemanjinom brief-u od 12.10 i odobrenom
rasporedu od osam slajdova. Jedna masina (krug) umesto tri stuba, otvara se
dokazom (rejz), jasno kaze sta je on-chain na Solani.
Stil kao v6-v10 (jedna ideja po slajdu, recenica u naslovu), tamna pozadina,
zuti naslovi. Mesta za prave fotografije su isprekidani okviri.
Tekst je i u content/colosseum/deck.md (izvor istine za tekst i izvore);
kad se menja tekst, menjaju se oba.
Pokretanje: python3 infra/deck_build_v11.py <izlaz.pptx>
Ikonice: Noto Emoji v2.042 (Google, Apache 2.0), u infra/deck_icons/."""
import sys, os
ICONS=os.path.join(os.path.dirname(os.path.abspath(__file__)),'deck_icons')
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
from lxml import etree

Y=RGBColor(0xEF,0xB7,0x3E); W=RGBColor(0xFF,0xFF,0xFF); G=RGBColor(0x9A,0x9A,0xA0)
L=RGBColor(0x8A,0x8A,0x90); D=RGBColor(0x2A,0x2A,0x2E); GR=RGBColor(0x6F,0xC2,0x8A); BL=RGBColor(0x7F,0xB3,0xE6)
FONT='Poppins'
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
C,Lf,Rt=PP_ALIGN.CENTER,PP_ALIGN.LEFT,PP_ALIGN.RIGHT

def slide(notes=''):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    f=s.background.fill; f.gradient(); f.gradient_angle=315
    f.gradient_stops[0].color.rgb=RGBColor(0x1C,0x1C,0x1F); f.gradient_stops[1].color.rgb=RGBColor(0x38,0x37,0x3C)
    if notes: s.notes_slide.notes_text_frame.text=notes
    return s

def tb(s,x,y,w,h,runs,size=20,color=W,bold=False,align=C,anchor=MSO_ANCHOR.TOP,font=FONT):
    t=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)).text_frame
    t.word_wrap=True; t.vertical_anchor=anchor
    if isinstance(runs,str): runs=[(runs,color)]
    p=t.paragraphs[0]; p.alignment=align
    for text,col in runs:
        parts=text.split('\n')
        for i,part in enumerate(parts):
            if i>0: p=t.add_paragraph(); p.alignment=align
            if not part: continue
            r=p.add_run(); r.text=part; r.font.size=Pt(size); r.font.bold=bold
            r.font.color.rgb=col
            if font: r.font.name=font
    return t

def icon_pic(s,name,cx,y,size=1.25):
    s.shapes.add_picture(os.path.join(ICONS,name+'.png'),Inches(cx-size/2),Inches(y),Inches(size),Inches(size))

def arrow(s,x1,y,x2,color=L):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y),Inches(x2),Inches(y))
    c.line.color.rgb=color; c.line.width=Pt(1.5)
    ln=c.line._get_or_add_ln(); ln.append(etree.SubElement(ln,qn('a:tailEnd'),type='triangle'))

def rect(s,x,y,w,h,fill=D,line=L,lw=1.0,rounded=True,dash=False):
    r=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    if fill is None: r.fill.background()
    else: r.fill.solid(); r.fill.fore_color.rgb=fill
    if line is None: r.line.fill.background()
    else:
        r.line.color.rgb=line; r.line.width=Pt(lw)
        if dash: r.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    r.shadow.inherit=False
    return r

def photo(s,x,y,w,h,label):
    """Mesto za pravu fotografiju: isprekidan okvir sa opisom. Zamenjuje se u Slides-u."""
    rect(s,x,y,w,h,D,L,1.25,False,True)
    tb(s,x+0.2,y+h/2-0.5,w-0.4,1.0,label,14,G,False,C,MSO_ANCHOR.MIDDLE)

def head1(s,text,size=34,w=11.7): tb(s,0.8,0.5,w,1.6,text,size,Y,True,Lf)
def sub(s,text,y=2.2,size=20,w=11.7,color=W): tb(s,0.8,y,w,0.9,text,size,color,False,Lf)
def foot(s,text,color=G,size=13): tb(s,0.8,6.85,11.7,0.4,text,size,color,False,Lf)
def tag(s,x,y,text,color):
    r=rect(s,x,y,1.4,0.34,None,color,1.0)
    tb(s,x,y+0.02,1.4,0.3,text,10,color,True,C)

def line_arrow(s,x1,y1,x2,y2,color=Y,width=2.0):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    c.line.color.rgb=color; c.line.width=Pt(width)
    ln=c.line._get_or_add_ln(); ln.append(etree.SubElement(ln,qn('a:tailEnd'),type='triangle'))

# 1 cover
s=slide('Hi, I am Nemanja. We are HiveBits. We build community-owned farms and we sell what they produce. The first farm is already funded.\n\nA: $140K raised, 150 people, 27 hours (Nemanja 12.10; MetaDAO screenshot, news.md 10.10; post @hivebits_io 08.09.2026). Main sentence is the approved one from 12.10, word for word. C: "builds farms" in plural and "sells" are the model; one farm is funded and under construction, sales have not started (said on slides 4 and 5).\n\nFOTO: desno je mesto za pravu fotografiju HiveBits kosnica (Drive hivebits/Swarm). Ne AI slika.')
tb(s,0.8,0.7,6.6,0.9,'HiveBits',40,W,True,Lf)
tb(s,0.8,1.9,6.6,2.4,'HiveBits builds community-owned farms and sells what they produce.',34,Y,True,Lf)
tb(s,0.8,4.6,6.6,1.2,'The first farm is already funded: $140K from 150 people in 27 hours.',20,W,False,Lf)
photo(s,7.7,0.6,4.9,6.3,'Photo: real HiveBits hives\n(from Drive)')

# 2 proof
s=slide('[Proof]\n\nWe raised 140 thousand dollars through our token, NECTAR, on Futard, in the MetaDAO ecosystem. 150 people funded 300 smart hives in 27 hours. It is the first tokenized farm owned by the community. The raise and the token are on Solana. The farm is being built in Bosnia right now, and we film every step.\n\nA: "we raised $140K through $NECTAR on Futard for the first tokenized farm owned by the community" (Nemanja 12.10). 150 funders, $140K target, $243K offered, 173% (MetaDAO screenshot, news.md 10.10). 27 hours: post @hivebits_io 08.09.2026 ("one day and three hours"). futard.io/launch/AbTuTsoBJS2jQggvfb7d1LA2SfbHa5NB5HbJqX8C2PnS. "First tokenized smart bee farm is live on @solana": @SuperteamBR, RT @hivebits_io 07.09.2026. B: farm under construction in Bosnia, documented (Nemanja brief 11.10; post @hivebits_io 16.09). On-chain today: the raise and the token. Nothing else on this slide is claimed as on-chain.\n\nNE PISE: "Target: $140,000", $243K (#49), revenue, returns.\n\nSLIKE: levo screenshot Futard stranice (Nemanja ga daje), desno prava fotka farme.')
head1(s,'We raised $140K for the first tokenized farm owned by the community.',32)
for i,(num,lab) in enumerate([('150','people'),('300','smart hives'),('27','hours')]):
    x=0.8+i*2.2
    tb(s,x,2.05,2.1,0.7,num,40,Y,True,Lf)
    tb(s,x,2.8,2.1,0.4,lab,16,W,False,Lf)
rect(s,7.3,1.95,5.3,1.4,D,GR,1.25)
tb(s,7.5,2.0,4.9,0.4,'On Solana: the raise and the token.',15,GR,True,Lf)
tb(s,7.5,2.5,4.9,0.8,'Raised through $NECTAR on Futard, in the MetaDAO ecosystem.',13,W,False,Lf)
tb(s,0.8,3.45,11.7,0.5,'The farm is being built in Bosnia. We record every step.',19,W,True,Lf)
photo(s,0.8,4.1,5.7,2.7,'Screenshot: Futard launch page\n(provided by the founder)')
photo(s,6.9,4.1,5.7,2.7,'Photo: the farm in Bosnia\n(real)')

# 3 problem (slajd iz v10, tekst nepromenjen)
s=slide('[Problem]\n\nWhy do we start from the farm? Because the farmer does not set his price. In the US, a beekeeper gets $2.45 a pound from a packer. If he sells the same honey himself, $7.15.\n\nA: USDA NASS Honey, March 2026: 2025 average price sold to co-ops/packers $2.45/lb; retail (beekeeper sells direct) $7.15/lb. https://esmis.nal.usda.gov/sites/default/release-files/795818/hony0326.txt (market.md 4a). One category, one country, one primary source. C: "shortens the distance" is intent. Tekst na slajdu je isti kao v10 slajd 2.\n\nNE GOVORI: "all intermediaries are unnecessary", "every farmer gets one third".')
head1(s,'Agriculture needs a better connection between production and customers.',32)
sub(s,'Production, financing, distribution and retail are separate worlds. The farmer meets only the first one.',2.2,19)
yc=3.5
nodes=[(1.7,'farmer','Beekeeper'),(5.0,'factory','Packer'),(8.3,'store','Store'),(11.6,'customer','Customer')]
for x,icon,label in nodes:
    icon_pic(s,icon,x,yc,1.1); tb(s,x-1.3,yc+1.15,2.6,0.4,label,16,W,True)
for i in range(3): arrow(s,nodes[i][0]+0.75,yc+0.55,nodes[i+1][0]-0.75)
tb(s,0.4,yc+1.6,2.6,0.6,'$2.45 / lb',26,Y,True)
tb(s,10.3,yc+1.6,2.6,0.6,'$7.15 / lb',26,Y,True)
tb(s,3.0,yc+1.7,7.3,0.5,'US honey, 2025. Same beekeeper: to a packer, or sold direct.',15,G,False)
tb(s,0.8,6.1,11.7,0.6,'HiveBits shortens the distance between the farm and the people who eat what it makes.',20,W,True,Lf)
foot(s,'Source: USDA NASS Honey report, March 2026, 2025 average prices.')

# 4 marketplace
s=slide('[Marketplace]\n\nNow we are building the marketplace. You can rent a whole hive or a cow for a season, or you can just buy the product. We sell our own products first. Then we invite other beekeepers, dairy farmers and other farmers. We take 5 percent of every sale. We start with farmers we know. Our hive sensors show weight, temperature and humidity. They show what happens. They do not certify quality.\n\nC: marketplace is being built; no sales and no rentals yet (Nemanja 10.10 and 12.10). Rent a whole hive or a cow or just buy the product; our products first, then other beekeepers, dairy farmers and others; 5% of every sale: Nemanja 12.10, word for word. B: hive sensors measure weight, temperature, humidity (Nemanja brief 11.10; post @hivebits_io "every colony monitored by our own hardware"). Start with farmers we know: Nemanja brief 11.10.\n\nNE PISE: escrow, payment method, lab test, fee on rental as a separate number, sensors on cows.\n\nFOTO: desno mesto za pravu fotografiju pametne kosnice sa senzorom.')
head1(s,'Now we are building the marketplace that sells what farms produce.',32)
for i,(t,body) in enumerate([('RENT','Rent a whole hive or a cow for a season.'),('BUY','Or just buy the product.')]):
    x=0.8+i*3.5
    rect(s,x,2.2,3.3,1.55,D,Y,1.25)
    tb(s,x+0.25,2.3,2.8,0.45,t,20,Y,True,Lf)
    tb(s,x+0.25,2.8,2.85,0.9,body,16,W,False,Lf)
rows=[('Who sells','Our products first. Then other beekeepers, dairy farmers and other farmers.'),
      ('How we earn','5% of every sale.'),
      ('Trust','We start with farmers we know. Our hive sensors show weight, temperature and humidity. A sensor shows what happens. It does not certify quality.')]
ys=[4.0,4.9,5.45]
for (lab,txt),y in zip(rows,ys):
    tb(s,0.8,y,1.6,0.4,lab,14,Y,True,Lf)
    tb(s,2.4,y-0.03,5.2,1.3,txt,16,W,False,Lf)
photo(s,8.1,2.2,4.5,4.5,'Photo: smart hive with sensor\n(real)')

# 5 loop
s=slide('[The loop]\n\nThis is one machine, not three products. The community funds a farm on-chain. The farm produces, and the sensors show it. The products sell on the marketplace. Sales show that the farm works, so the community funds the next one. The new farm brings new products and new buyers. The people who fund a farm are also its first buyers. So every raise brings supply and demand. The first step is done. The marketplace closes the loop.\n\nA: step 1 happened once ($140K, 300 hives, Futard, on Solana). B: sensors exist on our hives. C: steps 3, 4 and 5 are the model (Nemanja 12.10: the community invests in new farms, products are sold through the same marketplace). "Funders are the first buyers" is the design of the model, not a measured result: no sale has happened yet. Do not state it as a number or a conversion rate. Only confirmed revenue: 5% of every sale. A fee on funding new farms and hardware sales are NOT confirmed and are not on the slide.')
head1(s,'One loop: the community funds a farm, the farm produces, the marketplace sells.',30)
bw,bh=3.6,1.5; xs=[0.8,4.87,8.93]; yt,yb=2.0,4.05
def node(x,y,num,icon,text,line=L,lw=1.0):
    rect(s,x,y,bw,bh,D,line,lw)
    icon_pic(s,icon,x+0.55,y+0.38,0.7)
    tb(s,x+1.05,y+0.08,0.4,0.4,num,13,Y,True,Lf)
    tb(s,x+1.05,y+0.38,bw-1.15,bh-0.4,text,14,W,False,Lf)
node(xs[0],yt,'1','coin','The community funds a farm, on-chain.',GR,1.75)
node(xs[1],yt,'2','bee','The farm produces. Sensors show it.')
node(xs[2],yt,'3','store','Products sell on the marketplace. 5% to HiveBits.')
node(xs[2],yb,'4','handshake','Sales show the farm works. The community funds the next one.')
node(xs[0],yb,'5','seedling','The new farm brings new products and new buyers.')
line_arrow(s,xs[0]+bw+0.05,yt+bh/2,xs[1]-0.05,yt+bh/2)
line_arrow(s,xs[1]+bw+0.05,yt+bh/2,xs[2]-0.05,yt+bh/2)
line_arrow(s,xs[2]+bw/2,yt+bh+0.05,xs[2]+bw/2,yb-0.05)
line_arrow(s,xs[2]-0.05,yb+bh/2,xs[0]+bw+0.05,yb+bh/2)
line_arrow(s,xs[0]+bw/2,yb-0.05,xs[0]+bw/2,yt+bh+0.05)
tb(s,xs[1],yb+0.05,bw,0.6,'the next farm',13,G,False,C)
tb(s,0.8,5.75,11.7,0.9,'The people who fund a farm are also its first buyers. Every raise brings supply and demand.',18,Y,True,Lf)
foot(s,'We did step 1 with our own farm. The marketplace closes the loop, and that is what we are building now.',W,14)

# 6 market and expansion
s=slide('[Market and expansion]\n\nWe start with honey. In Texas, households buy about 71 million dollars of honey a year. In the European Union it is between 2.6 and 5.2 billion euros. Our vision is bigger than honey: new farms funded by the community, other kinds of sustainable agriculture, and one marketplace for all of it.\n\nMETHOD: households (or tonnes eaten) x yearly honey spend (or retail price) = size of the honey category per market. This is category demand, not our revenue. SAM and SOM are not calculated (#53).\nTexas: 10,992,816 households 2020-24 (US Census QuickFacts, https://www.census.gov/quickfacts/fact/csv/TX/HSD410224). 35.7% of households buy honey, $18.15 a year per buying household: NielsenIQ panel for National Honey Board, 52 wks to 2 Nov 2024, whole US (https://honey.com/images/files/National-Honey-Board-Category-Review-Analysis-2024.pdf). ASSUMPTION: national share and spend applied to Texas. 10.99M x 0.357 x 18.15 = $71.2M.\nEU: 436.5K t apparent consumption (2022, CBI, https://www.cbi.eu/market-information/honey/certified-honey/market-potential); retail EUR 5.58-24.94/kg (Stiftung Warentest 4/2025, Germany, 24 honeys), EUR 6-12 taken as the typical range. Result is a range.\nMontreal (about 4,200 t a year, no retail price from a primary source) and Brazil (about 12,000 t a year, supply market first) are off the slide; numbers and sources stay in deck.md.\nC: expansion is vision (Nemanja 12.10: the community invests in new farms and other kinds of sustainable agriculture, products sold through the same marketplace). No numbers for other categories.')
head1(s,'We start with honey. The loop is built for other farms too.',32)
tb(s,0.8,1.9,7.4,0.4,'Honey bought per year, bottom-up',15,G,True,Lf)
mk=[('Texas','11.0M households\nx 35.7% buy honey\nx $18.15 a year','about $71M a year'),
    ('European Union','436K t eaten a year\nx EUR 6 to 12 per kg retail','EUR 2.6 to 5.2B a year')]
for i,(name,calc,res) in enumerate(mk):
    x=0.8+i*3.75
    rect(s,x,2.4,3.55,3.6,D,Y,1.25)
    tb(s,x+0.25,2.55,3.05,0.5,name,22,W,True,Lf)
    tb(s,x+0.25,3.2,3.05,1.4,calc,15,G,False,Lf)
    tb(s,x+0.25,4.75,3.05,1.1,res,22,Y,True,Lf)
tb(s,8.75,1.9,3.8,0.4,'Vision',15,Y,True,Lf)
vis=['New farms, funded by the community.','Other kinds of sustainable agriculture.','One marketplace sells it all.']
for i,v in enumerate(vis):
    y=2.4+i*1.25
    rect(s,8.75,y,3.85,1.0,D,L,1.0)
    tb(s,8.95,y,3.45,1.0,v,16,W,False,Lf,MSO_ANCHOR.MIDDLE)
    if i<2: line_arrow(s,10.67,y+1.03,10.67,y+1.22,L,1.5)
foot(s,'Sources: US Census, NielsenIQ for National Honey Board (2024), CBI (2022), Stiftung Warentest (2025). Method in notes.',G,12)

# 7 competition
s=slide('[Competition]\n\nOthers do parts of this. Pollenity does hive adoption. CrowdFarming lets you adopt a tree. AgriDex trades commodities on Solana. AcreTrader and FarmTogether sell farmland. Each of them holds one part of the loop. As far as we found, nobody holds the whole loop.\n\nIZVORI (web pretraga 11.10): Pollenity: EBRD news 2026 (Adopt-a-Hive subscription, single-origin honey), CORDIS 887791 (BeeBot, HiHive). CrowdFarming: Euronews 05.06.2026 (adoptions, 300+ producers), FreshPlaza; fee 16-32% (market.md). AgriDex: The Defiant ($9M stablecoin trades, B2B, Africa), The Block (first RWA trade on Solana, 0.15% per side). AcreTrader: accredited only, $10-25K minimum (2026 review). FarmTogether: accredited, $15K crowdfunding minimum, $50K fund (farmtogether.com). Competitor metrics stay in notes, not on the slide. "As far as we found": competition.md 24.09 (Colosseum corpus, other chains) plus web search 11.10; Frontier 2026 corpus not searched. Drenio is off the slide: not verifiable (drenio.rs does not load, web search 11.10 and 12.10 finds nothing), #52.\n\nHiveBits row: filled = exists today (community funding done once with our farm; sensors on our hives). Outline = in build (rent, buy, beyond one category).')
head1(s,'Each of them holds one part of the loop. As far as we found, none holds the whole loop.',28)
colsx=[4.6,6.2,7.8,9.4,11.0]; heads=['Fund\nproduction','Rent a hive\nor animal','Buy\ndirect','Farm\nsensors','Beyond one\ncategory']
for x,h in zip(colsx,heads): tb(s,x-0.75,1.95,1.5,0.7,h,12,G,True,C)
rows=[('Pollenity','smart hives, hive adoption',['-','y','y','y','-']),
      ('CrowdFarming','adopt a tree or animal',['-','y','y','-','y']),
      ('AgriDex','B2B trade on Solana',['-','-','b','-','y']),
      ('AcreTrader','farmland investment',['y','-','-','-','-']),
      ('FarmTogether','farmland investment',['y','-','-','-','-']),
      ('HiveBits','farms plus marketplace',['h','p','p','h','p'])]
y0=2.75; rh=0.58
for i,(name,desc,cells) in enumerate(rows):
    y=y0+i*rh
    if name=='HiveBits': rect(s,0.7,y-0.03,11.3,rh,RGBColor(0x3A,0x36,0x28),Y,1.0)
    tb(s,0.8,y,2.0,rh,name,14,Y if name=='HiveBits' else W,True,Lf,MSO_ANCHOR.MIDDLE)
    tb(s,2.75,y,1.6,rh,desc,10,G,False,Lf,MSO_ANCHOR.MIDDLE)
    for x,c in zip(colsx,cells):
        cy=y+rh/2-0.14
        if c=='y': rect(s,x-0.14,cy,0.28,0.28,W,None)
        elif c=='h': rect(s,x-0.14,cy,0.28,0.28,Y,None)
        elif c=='p': rect(s,x-0.14,cy,0.28,0.28,None,Y,1.5)
        elif c=='b': rect(s,x-0.14,cy,0.28,0.28,W,None); tb(s,x-0.6,cy+0.26,1.2,0.25,'B2B only',9,G,False,C)
        else: tb(s,x-0.3,cy-0.06,0.6,0.3,'-',14,L,False,C)
ly=6.35
rect(s,0.8,ly+0.05,0.22,0.22,W,None); tb(s,1.1,ly,2.2,0.35,'offers today',11,G,False,Lf)
rect(s,3.0,ly+0.05,0.22,0.22,Y,None); tb(s,3.3,ly,2.4,0.35,'HiveBits, today',11,G,False,Lf)
rect(s,5.4,ly+0.05,0.22,0.22,None,Y,1.5); tb(s,5.7,ly,3.0,0.35,'HiveBits, in build',11,G,False,Lf)
foot(s,'All of these work. HiveBits puts the parts in one loop.',W,14)

# 8 team and next
s=slide('[Team and next]\n\nI am a beekeeper and an engineer. Sinisa manages the farm. We were at Colosseum twice with the smart hive. Since then we raised the money and started to build the farm. Next we finish the farm, launch the marketplace, and fund the next farm with the community. Thank you.\n\nB: team, two confirmed sentences (Nemanja brief 11.10). Eduardo is not in the team (Nemanja 10.10). A: two earlier Colosseum entries, Cypherpunk 2025 (https://colosseum.com/projects/explore/hivebits) and Frontier 2026 (https://colosseum.com/projects/explore/hivebits-1), both pitched the smart hive, no prize. This is the disclosure of earlier work the FAQ asks for. C: five next steps, all from Nemanja 12.10, no dates, no amounts. No ask (not confirmed, #51). GitHub status not verified (#22).\n\nFOTO: levo dole, osnivaci.')
head1(s,'We run real hives. Now we build the marketplace.',34)
tb(s,0.8,1.9,5.0,0.4,'Team',16,Y,True,Lf)
tb(s,0.8,2.3,5.0,1.9,[('Nemanja Šćepanović\n',W),('Beekeeper and engineer. Runs an apitherapy centre.\n\n',G),('Siniša Kežić\n',W),('Manages the farm.',G)],15,W,False,Lf)
photo(s,0.8,4.3,5.0,2.4,'Photo: founders (real)')
tb(s,6.5,1.9,6.1,0.4,'Next',16,Y,True,Lf)
tb(s,6.5,2.3,6.1,2.6,'1.  Finish the farm in Bosnia. Keep recording every step.\n2.  Launch the marketplace with our own products.\n3.  Invite the first beekeepers and dairy farmers.\n4.  Open renting: a whole hive for a season.\n5.  Fund the next farm with the community.',16,W,False,Lf)
rect(s,6.5,5.0,6.1,1.7,D,L,1.0)
tb(s,6.75,5.08,5.6,0.4,'Colosseum before',13,Y,True,Lf)
tb(s,6.75,5.45,5.6,1.2,'Cypherpunk 2025 and Frontier 2026. Both pitched the smart hive. Since then: the raise, the farm, and now the marketplace.',14,W,False,Lf)

prs.save(sys.argv[1]); print(len(prs.slides._sldIdLst),'slides')
