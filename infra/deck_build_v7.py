"""Colosseum deck v7 (2026-10-10): platforma za zajednicko vlasnistvo nad proizvodnjom.
Poceli kao community owned farm, platforma to ponavlja za druge farme. Stil kao v6 (YC: jedna
ideja po slajdu, recenica u naslovu; redosled po Pixar strukturi), tamna pozadina, zuti naslovi.
Tekst je i u content/colosseum/deck.md (izvor istine za tekst i izvore).
Pokretanje: python3 infra/deck_build_v7.py <izlaz.pptx>
Ikonice: Noto Emoji v2.042 (Google, Apache 2.0), u infra/deck_icons/."""
import sys, math, os
ICONS=os.path.join(os.path.dirname(os.path.abspath(__file__)),'deck_icons')
def icon_pic(s,name,cx,y,size=1.25):
    s.shapes.add_picture(os.path.join(ICONS,name+'.png'),Inches(cx-size/2),Inches(y),Inches(size),Inches(size))
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

Y=RGBColor(0xEF,0xB7,0x3E); W=RGBColor(0xFF,0xFF,0xFF); G=RGBColor(0x9A,0x9A,0xA0)
R=RGBColor(0xD9,0x60,0x5A); L=RGBColor(0x8A,0x8A,0x90)
FONT='Poppins'
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
C,Lf,Rt=PP_ALIGN.CENTER,PP_ALIGN.LEFT,PP_ALIGN.RIGHT

def slide(notes=''):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    f=s.background.fill; f.gradient(); f.gradient_angle=315
    f.gradient_stops[0].color.rgb=RGBColor(0x1C,0x1C,0x1F); f.gradient_stops[1].color.rgb=RGBColor(0x38,0x37,0x3C)
    if notes: s.notes_slide.notes_text_frame.text=notes
    return s

def tb(s,x,y,w,h,runs,size=20,color=W,bold=False,align=C,anchor=MSO_ANCHOR.TOP,font=FONT):
    """runs: string, ili lista (tekst, boja) za vise boja u jednom redu; '\n' pravi novi pasus."""
    t=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)).text_frame
    t.word_wrap=True; t.vertical_anchor=anchor
    if isinstance(runs,str): runs=[(runs,color)]
    p=t.paragraphs[0]; p.alignment=align
    for text,col in runs:
        parts=text.split('\n')
        for i,part in enumerate(parts):
            if i>0: p=t.add_paragraph(); p.alignment=align
            if not part: continue
            r=p.add_run(); r.text=part; r.font.size=Pt(size); r.font.bold=bold
            r.font.color.rgb=col
            if font: r.font.name=font
    return t

def arrow(s,x1,y,x2):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y),Inches(x2),Inches(y))
    c.line.color.rgb=L; c.line.width=Pt(1.25)
    ln=c.line._get_or_add_ln(); ln.append(etree.SubElement(ln,qn('a:tailEnd'),type='triangle'))

NODES=[(1.7,'farmer','Farmer'),(5.0,'factory','Food company'),(8.3,'store','Store'),(11.6,'customer','Customer')]
def chain(s,y=1.5):
    for x,icon,label in NODES:
        icon_pic(s,icon,x,y)
        tb(s,x-1.4,y+1.35,2.8,0.5,label,18,W,True)
    for i in range(3):
        arrow(s,NODES[i][0]+0.9,y+0.7,NODES[i+1][0]-0.9)

def head1(s,text,size=38): tb(s,0.8,0.6,11.7,1.6,text,size,Y,True,Lf)
def big(s,y,text,size=30,color=W,bold=True,x=0.8,w=11.7,h=0.9): tb(s,x,y,w,h,text,size,color,bold,Lf)
def xs(s,y=1.5):
    for i in range(3):
        mid=(NODES[i][0]+NODES[i+1][0])/2
        tb(s,mid-0.4,y+0.33,0.8,0.7,'X',30,R,True)

# 1 cover
s=slide('We are HiveBits. We started as the first community owned farm. We asked for $140K, and in 27 hours 150 people offered $243K. '
        'Now we found a way to repeat that for other farms.\n\n'
        'IZVOR: HiveBits na MetaDAO launchpadu, $NECTAR, status COMPLETE (screenshot, news.md 10.10.2026): cilj $140K, ponudjeno $243K, 173%, 150 funders. '
        '"The first community owned farm" su Nemanjine reci; "first" nije provereno.')
tb(s,0.8,1.0,11.7,1.3,'HiveBits',66,W,True,Lf)
big(s,2.6,'We started as the first community owned farm.',34,Y)
big(s,3.5,'We found a way to repeat it.',34,Y)
big(s,4.9,'150 people offered $243K in 27 hours.',32)

# 2 once upon a time
s=slide('[Once upon a time]\n\nHave you heard of the term "safe food"? It is food where you can check where it comes from and what is inside.\n\n'
        'Most people know only two kinds. Conventional food and organic food.\n\nDo you know what you eat?')
head1(s,'Do you know what you eat?',44)
tb(s,0.8,3.3,3.9,1.6,'Conventional\nfood',32,W,True)
tb(s,4.72,3.3,3.9,1.6,'Safe food',32,Y,True)
tb(s,8.63,3.3,3.9,1.6,'Organic food',32,W,True)

# 3 and every day
s=slide('[And every day]\n\nBetween you and the farmer there are two companies. You never see him. Take milk in Bosnia. You pay $1.09 for a liter. The farmer gets $0.34.\n\n'
        'He does not set his price. The dairy sets it, because the dairy is the only one who pays him. '
        'When cheap milk comes from import, the dairy pays him less. This year his price went down 25 percent, and he has nobody else to sell to.\n\n'
        'IZVOR: otkupna cena Viteski.ba 29.09.2026 ($0.34, pad 25%), cena na polici 072info 29.07.2026 ($1.09). Dva izvora, dva datuma. '
        'Sirovo mleko naspram preradjenog na polici: deo razlike je prerada i dostava, zato se govori "od $1.09 farmeru stigne $0.34", ne "mlekara uzima dve trecine". '
        'Srednja cena ($0.73) je nasa procena i nije na slajdu.')
head1(s,'You never see the farmer. And he does not set his price.',34)
Y3=2.3
chain(s,Y3)
tb(s,NODES[0][0]-1.4,Y3+1.9,2.8,0.7,'$0.34',34,Y,True)
tb(s,NODES[3][0]-1.4,Y3+1.9,2.8,0.7,'$1.09',34,Y,True)
tb(s,3.2,Y3+2.0,6.9,0.6,'one liter of milk in Bosnia',22,W,False)
tb(s,0.8,5.9,11.7,0.8,[('His price is ',W),('25% lower',R),(' than last year.',W)],28,W,True)

# 4 until one day: solution
s=slide('[Until one day]\n\nWe did it the other way. We did not go to a dairy or a bank. We let anyone invest in our farm and own it.\n\n'
        '150 people did. They own the farm with 300 smart hives. What the farm makes goes to them, on Solana. They know exactly what they eat, because they own the farm.\n\n'
        'HiveBits lets anyone do this with other farms. That was the idea from day one: anyone can invest in food production.\n\n'
        'IZVOR: 300 kosnica, $140K, MetaDAO: javni postovi @hivebits_io 07.09 i 08.09.2026; 150 ulagaca: screenshot MetaDAO (news.md 10.10). '
        '"They own the farm" i "oni su vlasnici, mi im ne prodajemo": Nemanja 10.10; prihod ide u trezor vlasnika tokena na lancu (javni post @hivebits_io 07.09).')
head1(s,'Anyone can invest in a farm and own it.',38)
Y4=2.1
chain(s,Y4); xs(s,Y4)
x1,x2,y0,depth=NODES[0][0],NODES[3][0],Y4+2.55,1.3
pts=[]
for k in range(0,61):
    a=math.pi*k/60
    pts.append((Inches((x1+x2)/2-(x2-x1)/2*math.cos(a)),Inches(y0+depth*math.sin(a))))
fb=s.shapes.build_freeform(pts[0][0],pts[0][1]); fb.add_line_segments(pts[1:],close=False)
arc=fb.convert_to_shape(); arc.fill.background(); arc.line.color.rgb=Y; arc.line.width=Pt(2.5)
ln=arc.line._get_or_add_ln(); etree.SubElement(ln,qn('a:tailEnd'),type='triangle')
tb(s,4.17,Y4+2.55,5.0,0.6,'HiveBits',30,Y,True)
tb(s,3.67,Y4+3.1,6.0,0.5,'We did this with our own farm first.',22,W,True)

# 5 the platform: three ways
s=slide('[Until one day]\n\nOn HiveBits you pay the farmer directly, in three ways.\n\n'
        'You invest in the farm and own a part of it. That is what 150 people did with ours.\n\n'
        'You rent a hive or a cow for a season, and you get what it makes.\n\n'
        'Or you simply buy the product, like in any shop.\n\n'
        'NAPOMENA: korak 1 je dokazan (nasa farma). Korak 2 i 3 jos nemaju prodaju na platformi; to se ne krije, slajd 7 kaze sta je dokazano.')
head1(s,'One platform. Own, rent, or buy.')
big(s,2.4,'1.  Invest in the farm and own a part of it.',38)
big(s,3.6,'2.  Rent a hive or a cow for a season.',38)
big(s,4.8,'3.  Buy the product.',38)
big(s,5.9,'We did step 1 first, with our own farm.',28,Y)

# 6 and because of that: why trust
s=slide('[And because of that]\n\nWhy would you trust a farm you never saw?\n\n'
        'First, we check the farm and license it. We are farmers, we know what to look for. We do not just collect the money.\n\n'
        'Second, the ownership and the money are on Solana. Everyone can see who owns the farm and where the revenue goes.\n\n'
        'On our own farm we go one step further. Every hive has a sensor, and you see the weight, the temperature and the humidity live.\n\n'
        'IZVOR: provera i licenciranje farmi: Nemanja 10.10 ("ne samo da se prikuplja novac"). Vlasnistvo i prihod na lancu: javni post @hivebits_io 07.09 (onchain-governed, tokenholder-owned treasury). '
        'Senzori samo za nasu farmu (Nemanja 10.10: za druge farme ne znamo kako): x.com/hivebits_io/status/2096269121751220404. '
        'Novac zakljucan do isporuke vazi samo za rentanje i kupovinu, ne za ulaganje (Nemanja 10.10), zato nije na slajdu; ide u dodatak A1.')
head1(s,'We check the farm. Solana shows who owns it and where the money goes.',32)
big(s,2.5,'We check and license every farm.',34)
big(s,3.7,'Ownership and revenue are on Solana.',34)
big(s,4.9,'Our farm goes further: every hive has a live sensor.',34)

# 7 and because of that: traction
s=slide('[And because of that]\n\nWe asked for $140K to build a farm with 300 smart hives. In 27 hours, 150 people offered $243K.\n\n'
        'We could take only $140K. The other $103K had no farm to go to. That is the demand for the platform. '
        'That is why we build it.\n\n'
        'IZVOR: HiveBits na MetaDAO launchpadu, screenshot (news.md 10.10.2026): cilj $140K, ponudjeno $243K, 173%, 150 funders. 27 sati: javni post @hivebits_io 08.09.2026.')
head1(s,'We asked for $140K. 150 people offered $243K.')
tb(s,0.8,2.4,11.7,2.4,'$243K',150,Y,True,Lf)
big(s,5.2,'in 27 hours, on Solana',32)
big(s,6.0,'$103K had no farm to go to. Yet.',32,Y)

# 8 until finally: market
s=slide('[Until finally]\n\nWe start with honey, because that is our farm. 46 million households in the US buy honey. They spend $839M a year. Our 5 percent of that is $42M, from honey alone.\n\n'
        'Milk is the same story, you saw the numbers. And the model does not care what the farm makes. '
        'Today it is farms. The same model works for anything that produces.\n\n'
        'IZVOR: NielsenIQ for National Honey Board, 52 weeks to 2 Nov 2024 (46.2M domacinstava x $18.15 = $839M), honey.com/images/files/National-Honey-Board-Category-Review-Analysis-2024.pdf. 5%: HiveBits (Nemanja 24.09). '
        'Za "vlasnistvo nad farmama" nema brojke sa izvorom, zato je sirina samo recenica.')
head1(s,'Honey first. Then milk. Then anything that produces.',34)
big(s,2.5,'46 million US households buy honey. $839M a year.',32)
big(s,3.7,'Our 5% of honey alone is $42M.',32,Y)
big(s,4.9,'The same model works for anything that produces.',32)

# 9 team and ask
s=slide('We are farmers and engineers. We are here to win and to join the Colosseum accelerator.\n\n'
        'Next is the second farm, funded the way we funded ours, and better ways to check the farms we bring.\n\n'
        'NAPOMENA: druga farma nije u razgovoru ni sa kim (Nemanja 10.10), zato je ona ask, ne traction.')
head1(s,'We are here to win and join the Colosseum accelerator.',34)
big(s,2.7,'Nemanja Šćepanović',32); big(s,3.4,'CEO, beekeeper. 10 years in engineering and IoT.',22,W,False)
big(s,4.2,'Siniša Kežić',32); big(s,4.9,'COO, chief beekeeper. Runs the farm build.',22,W,False)
big(s,5.9,'Next: the second farm.',28,Y,True,0.8,7.0)
big(s,5.9,'hivebits.io',28,Y,True,8.0,4.5)

# appendix
s=slide('Dodatak. Ovi slajdovi nisu deo videa od tri minuta. Sluze za pitanja i za prijavu.')
tb(s,0.8,3.0,11.7,1.2,'Appendix',54,W,True,Lf)

# A1 how we earn
s=slide('We earn in three places. A fee when a farm is funded. A share of every rental. And 5 percent of every sale.\n\n'
        'For renting and buying, the money is locked on Solana in stablecoins until the farmer delivers.\n\n'
        'IZVOR: 5% od prodaje: Nemanja 24.09. Naknada pri finansiranju i deo od rentanja: Nemanja 10.10, bez procenta (open-questions #39, #40).')
head1(s,'We earn when a farm is funded, rented, or sells.')
big(s,2.5,'A fee when a farm is funded.',34)
big(s,3.7,'A share of every rental.',34)
big(s,4.9,'5% of every sale.',34)

# A2 go to market
s=slide('We start with what we know. Our own farm is funded and owned by 150 people.\n\n'
        'Next is a second farm, with a farmer we know, funded the same way.\n\n'
        'For renting and buying we can onboard a farm without funding it first. Then we bring other farms.\n\n'
        'NAPOMENA: onboarding farme za rentanje bez finansiranja: Nemanja 10.10.')
head1(s,'First our farm. Then farmers we know. Then the rest.')
big(s,2.5,'1.  Our own farm. Funded and owned by 150 people.',30)
big(s,3.7,'2.  A second farm, with a farmer we know.',30)
big(s,4.9,'3.  Any farm can join for renting and buying.',30)

# A3 competition
s=slide('Many teams tried a farm marketplace here. None of them won. They were all software with no farmer behind it.\n\n'
        'We are farmers. We have our farm, our sensors, and 150 people already own a farm with us.\n\n'
        'IZVOR: Colosseum Copilot, cetiri hakatona od marta 2024 do septembra 2025, provereno 10.10.2026 (competition.md). 3Bee, Pollenity, Honeyverse rentaju samo svoje kosnice, van lanca.')
head1(s,'No farm marketplace won here. They had no farm.')
big(s,2.5,'We have our own farm.',34)
big(s,3.7,'We have our own hardware.',34)
big(s,4.9,'150 people already own a farm with us.',34)

prs.save(sys.argv[1]); print(len(prs.slides._sldIdLst),'slides')
