"""Colosseum deck v6 (YC: legible, simple, obvious; redosled po Pixar strukturi): tamna pozadina, zuti naslovi, po Nemanjinom Google Slides decku
("Colosseum Deck V1 - Marketplace"). Pokretanje: python3 infra/deck_build_v6.py <izlaz.pptx>
Ikonice: Noto Emoji v2.042 (Google, Apache 2.0), u infra/deck_icons/."""
import sys, math, os
ICONS=os.path.join(os.path.dirname(os.path.abspath(__file__)),'deck_icons')
def icon_pic(s,name,cx,y,size=1.25):
    s.shapes.add_picture(os.path.join(ICONS,name+'.png'),Inches(cx-size/2),Inches(y),Inches(size),Inches(size))
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

Y=RGBColor(0xEF,0xB7,0x3E); W=RGBColor(0xFF,0xFF,0xFF); G=RGBColor(0x9A,0x9A,0xA0)
R=RGBColor(0xD9,0x60,0x5A); L=RGBColor(0x8A,0x8A,0x90)
FONT='Poppins'
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
C,Lf,Rt=PP_ALIGN.CENTER,PP_ALIGN.LEFT,PP_ALIGN.RIGHT

def slide(notes=''):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    f=s.background.fill; f.gradient(); f.gradient_angle=315
    f.gradient_stops[0].color.rgb=RGBColor(0x1C,0x1C,0x1F); f.gradient_stops[1].color.rgb=RGBColor(0x38,0x37,0x3C)
    if notes: s.notes_slide.notes_text_frame.text=notes
    return s

def tb(s,x,y,w,h,runs,size=20,color=W,bold=False,align=C,anchor=MSO_ANCHOR.TOP,font=FONT):
    """runs: string, ili lista (tekst, boja) za vise boja u jednom redu; '\n' pravi novi pasus."""
    t=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)).text_frame
    t.word_wrap=True; t.vertical_anchor=anchor
    if isinstance(runs,str): runs=[(runs,color)]
    p=t.paragraphs[0]; p.alignment=align
    for text,col in runs:
        parts=text.split('\n')
        for i,part in enumerate(parts):
            if i>0: p=t.add_paragraph(); p.alignment=align
            if not part: continue
            r=p.add_run(); r.text=part; r.font.size=Pt(size); r.font.bold=bold
            r.font.color.rgb=col
            if font: r.font.name=font
    return t

def title(s,text): tb(s,0.8,0.55,11.7,0.9,text,32,Y,True,Lf)
def footer(s,text): tb(s,0.8,6.85,11.7,0.4,text,10,G,False,Lf)

def arrow(s,x1,y,x2):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y),Inches(x2),Inches(y))
    c.line.color.rgb=L; c.line.width=Pt(1.25)
    ln=c.line._get_or_add_ln(); ln.append(etree.SubElement(ln,qn('a:tailEnd'),type='triangle'))

NODES=[(1.7,'farmer','Farmer'),(5.0,'factory','Food company'),(8.3,'store','Store'),(11.6,'customer','Customer')]
def chain(s,y=1.5):
    for x,icon,label in NODES:
        icon_pic(s,icon,x,y)
        tb(s,x-1.4,y+1.35,2.8,0.5,label,18,W,True)
    for i in range(3):
        arrow(s,NODES[i][0]+0.9,y+0.7,NODES[i+1][0]-0.9)



# v6 pravila (Kevin Hale, YC): jedna ideja po slajdu, ideja pise u naslovu na vrhu,
# krupno i bold, bez sitnih napomena. Izvori i ograde idu u beleske, ne na slajd.
def head1(s,text,size=38): tb(s,0.8,0.6,11.7,1.6,text,size,Y,True,Lf)
def big(s,y,text,size=30,color=W,bold=True,x=0.8,w=11.7,h=0.9): tb(s,x,y,w,h,text,size,color,bold,Lf)
def xs(s,y=1.5):
    for i in range(3):
        mid=(NODES[i][0]+NODES[i+1][0])/2
        tb(s,mid-0.4,y+0.33,0.8,0.7,'X',30,R,True)

# 1 naslovni: ko smo i dokaz
s=slide('We are HiveBits. We build the first community owned farm. We asked for $140K, and in 27 hours 150 people offered $243K.\n\n'
        'IZVOR: HiveBits na MetaDAO launchpadu, $NECTAR. Cilj $140K, 173%, 150 funders.')
tb(s,0.8,1.2,11.7,1.3,'HiveBits',66,W,True,Lf)
big(s,2.9,'The first community owned farm.',36,Y)
big(s,4.3,'150 people offered $243K in 27 hours.',32)

# 2 Once upon a time: problem i uvid
s=slide('[Once upon a time]\n\nHave you heard of the term "safe food"? It is food where you can check where it comes from and what is inside.\n\n'
        'Most people know only two kinds. Conventional food and organic food.\n\nDo you know what you eat?')
head1(s,'Do you know what you eat?',44)
tb(s,0.8,3.3,3.9,1.6,'Conventional\nfood',32,W,True)
tb(s,4.72,3.3,3.9,1.6,'Safe food',32,Y,True)
tb(s,8.63,3.3,3.9,1.6,'Organic food',32,W,True)

# 3 And every day: zasto danasnji put ne radi
s=slide('[And every day]\n\nTake milk in Bosnia. The farmer gets only a third of the final price.\n\n'
        'When cheap milk comes from import, the dairy pays him less. This year his price went down 25 percent.\n\n'
        'The result is that the customer pays three times more than the farmer gets, for a processed product.\n\n'
        'IZVOR: otkupna cena Viteski.ba 29.09.2026, cena na polici 072info 29.07.2026. Mlekara prodaje prodavnici za oko $0.73, to je nasa procena i zato nije na slajdu.')
head1(s,'The farmer gets a third of what you pay.')
Y3=2.3
chain(s,Y3)
tb(s,NODES[0][0]-1.4,Y3+1.9,2.8,0.7,'$0.34',34,Y,True)
tb(s,NODES[3][0]-1.4,Y3+1.9,2.8,0.7,'$1.09',34,Y,True)
tb(s,3.2,Y3+2.0,6.9,0.6,'one liter of milk in Bosnia',22,W,False)
tb(s,0.8,5.9,11.7,0.8,[('His price is ',W),('25% lower',R),(' than last year.',W)],28,W,True)

# 4 Until one day: resenje
s=slide("[Until one day]\n\nPeople know that food from the farmer is better. But they don't know how to find a farmer they can trust. So they go to the store.\n\n"
        'HiveBits solves that. You buy food directly from farmers you can trust.\n\n'
        'For the same liter of milk, the farmer gets about 51 cents instead of 34. You pay about 75 instead of $1.09.\n\n'
        'NAPOMENA: $0.51 i $0.75 su nasa racunica (dostava dva puta nedeljno, 400 litara po turi), zato na slajdu pise "about".')
head1(s,'Buy food from trusted farmers.')
Y4=2.1
chain(s,Y4); xs(s,Y4)
x1,x2,y0,depth=NODES[0][0],NODES[3][0],Y4+2.55,1.3
pts=[]
for k in range(0,61):
    a=math.pi*k/60
    pts.append((Inches((x1+x2)/2-(x2-x1)/2*math.cos(a)),Inches(y0+depth*math.sin(a))))
fb=s.shapes.build_freeform(pts[0][0],pts[0][1]); fb.add_line_segments(pts[1:],close=False)
arc=fb.convert_to_shape(); arc.fill.background(); arc.line.color.rgb=Y; arc.line.width=Pt(2.5)
ln=arc.line._get_or_add_ln(); etree.SubElement(ln,qn('a:tailEnd'),type='triangle')
tb(s,4.17,Y4+2.7,5.0,0.6,'HiveBits',30,Y,True)
tb(s,3.67,Y4+3.3,6.0,0.5,'about $0.75 per liter',22,W,True)
tb(s,0.1,Y4+1.9,3.6,0.5,'about $0.51, not $0.34',17,Y,True,Lf)
tb(s,9.63,Y4+1.9,3.6,0.5,'about $0.75, not $1.09',17,Y,True,Rt)

# 5 Until one day, nastavak: sta tacno radi (lista koraka)
s=slide('[Until one day]\n\nWe already did the last step. People invested in our farm.\n\n'
        'The marketplace does the same for other farmers. We find the farm, and people can invest through us.\n\n'
        'Before that you can rent a hive or buy ahead. And you can simply buy the product, like in any shop.')
head1(s,'One marketplace, three steps.')
big(s,2.4,'1.  Buy.',40)
big(s,3.6,'2.  Rent or buy ahead.',40)
big(s,4.8,'3.  Invest in the farm.',40)
big(s,5.9,'We did step 3 first.',28,Y)

# 6 And because of that: zasto je bolje
s=slide('[And because of that]\n\nThe farmer gets more, because nobody stands between him and you.\n\n'
        'You know who made your food. On our farm every hive has a sensor. We track the weight, the temperature and the humidity.\n\n'
        'Your money is locked on Solana, in stablecoins. The farmer gets it when he delivers the food.')
head1(s,'Better for the farmer. Better for you.')
big(s,2.5,'The farmer gets more.',36)
big(s,3.7,'You know who made your food.',36)
big(s,4.9,'Your money is locked on Solana until delivery.',36)

# 7 And because of that: traction
s=slide('[And because of that]\n\nWe asked for $140K to build a farm with 300 smart hives. In 27 hours, 150 people offered $243K.\n\n'
        'We could take only $140K. The other $103K had no farm to go to. That is why we build the marketplace.\n\n'
        'IZVOR: HiveBits na MetaDAO launchpadu. Cilj $140K, ponudjeno $243K, 173%, 150 funders.')
head1(s,'We asked for $140K. 150 people offered $243K.')
tb(s,0.8,2.5,11.7,2.4,'$243K',150,Y,True,Lf)
big(s,5.4,'in 27 hours, on Solana',36)

# 8 Until finally: trziste
s=slide('[Until finally]\n\n46 million households in the US buy honey. They spend $839M a year. Our 5 percent of that is $42M, from honey alone.\n\n'
        'Later the same farmers sell to shops and restaurants. AgriDex already does farm trade on Solana, in bulk.\n\n'
        'IZVOR: NielsenIQ for National Honey Board, 52 weeks to 2 Nov 2024. 5%: HiveBits. AgriDex: The Defiant.')
head1(s,'Honey alone is a $42M market for us.')
big(s,2.5,'46 million US households buy honey.',34)
big(s,3.7,'They spend $839M a year.',34)
big(s,4.9,'Our 5% is $42M.',34,Y)

# 9 tim i ask
s=slide('We are beekeepers and engineers. We are here to win and to join the Colosseum accelerator.\n\n'
        'Next we bring more farmers and we build better ways to check them.')
head1(s,'We are here to win and join the Colosseum accelerator.',34)
big(s,2.9,'Nemanja Šćepanović',32); big(s,3.6,'CEO, beekeeper. 10 years in engineering and IoT.',22,W,False)
big(s,4.5,'Siniša Kežić',32); big(s,5.2,'COO, chief beekeeper. Runs the farm build.',22,W,False)
big(s,6.3,'hivebits.io',28,Y)

# ---- dodatak: za pitanja, ne za video ----
s=slide('Dodatak. Ovi slajdovi nisu deo videa od tri minuta. Sluze za pitanja i za prijavu.')
tb(s,0.8,3.0,11.7,1.2,'Appendix',54,W,True,Lf)

s=slide('We start with what we know. Our own farm with 300 smart hives is funded.\n\n'
        'Next is the second farm, with a beekeeper we know. It is funded the same way we funded ours.\n\n'
        'When this works for honey, we bring other farmers.')
head1(s,'First our farm. Then farmers we know.')
big(s,2.5,'1.  Our own farm. Funded in 27 hours.',30)
big(s,3.7,'2.  A second farm, with a beekeeper we know.',30)
big(s,4.9,'3.  Other farmers.',30)

s=slide('Many teams tried a farm marketplace here. None of them won. They were all software with no farmer behind it.\n\n'
        'We are farmers. We have our hives and our sensors, and 150 people already backed us.\n\n'
        'IZVOR: Colosseum Copilot, cetiri hakatona od marta 2024 do septembra 2025, provereno 10.10.2026. CrowdFarming 16 do 32 odsto: njihova stranica o transparentnosti. 3Bee, Pollenity, Honeyverse prodaju samo svoje kosnice, van lanca.')
head1(s,'No farm marketplace won here. They had no farm.')
big(s,2.5,'We have our own farm.',34)
big(s,3.7,'We have our own hardware.',34)
big(s,4.9,'We have 150 backers.',34)

s=slide('We earn in two places. A fee when a farm is funded, and 5 percent of every sale after that.\n\n'
        'CrowdFarming takes 16 to 32 percent.')
head1(s,'We earn when a farm is funded and on every sale.')
big(s,2.5,'A fee when a farm is funded.',34)
big(s,3.7,'5% of every sale.',34)

prs.save(sys.argv[1]); print(len(prs.slides._sldIdLst),'slides')
