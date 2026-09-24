import re,sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
md=open('content/colosseum/deck.md').read()
secs=re.split(r'\n## (\d+)\. ',md)[1:]
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
blank=prs.slide_layouts[6]
Y=RGBColor(0xF2,0xB1,0x00); K=RGBColor(0x1A,0x1A,0x1A); G=RGBColor(0x77,0x77,0x77)
def part(body,start,ends):
    m=re.search(r'\*\*'+start+r'\*\*\n(.*?)(?=\n\*\*(?:'+'|'.join(ends)+r')|\n---|\Z)',body,re.S)
    return m.group(1).strip() if m else ''
for n,body in zip(secs[0::2],secs[1::2]):
    title=body.split('\n',1)[0].strip()
    on=part(body,'Na slajdu',['Nemanja govori','Izvori'])
    talk=part(body,'Nemanja govori',['Izvori'])
    m=re.search(r'\*\*Izvori:?\*\*:?\s*(.*?)(?=\n---|\Z)',body,re.S); src=m.group(1).strip() if m else ''
    lines=[l.strip() for l in on.split('\n') if l.strip()]
    bullets=[re.sub(r'^- ','',l) for l in lines if l.startswith('- ')]
    footer=' '.join(l[len('Source:'):].strip() for l in lines if l.startswith('Source:'))
    if n=='1': title,bullets=bullets[0],bullets[1:]
    s=prs.slides.add_slide(blank)
    bar=s.shapes.add_shape(1,0,0,Inches(0.25),prs.slide_height); bar.fill.solid(); bar.fill.fore_color.rgb=Y; bar.line.fill.background()
    tb=s.shapes.add_textbox(Inches(0.8),Inches(0.5),Inches(11.8),Inches(1.2)).text_frame; tb.word_wrap=True
    p=tb.paragraphs[0]; p.text=title; p.font.size=Pt(36); p.font.bold=True; p.font.color.rgb=K
    bx=s.shapes.add_textbox(Inches(0.8),Inches(1.9),Inches(11.8),Inches(4.6)).text_frame; bx.word_wrap=True
    size=28 if len(bullets)<=3 else (24 if len(bullets)<=5 else 20)
    for i,b in enumerate(bullets):
        q=bx.paragraphs[0] if i==0 else bx.add_paragraph()
        q.text=b.replace('**',''); q.font.size=Pt(size); q.font.color.rgb=K; q.space_after=Pt(12 if size>20 else 8)
    if footer:
        fb=s.shapes.add_textbox(Inches(0.8),Inches(6.7),Inches(11.8),Inches(0.6)).text_frame; fb.word_wrap=True
        fp=fb.paragraphs[0]; fp.text='Source: '+footer; fp.font.size=Pt(11); fp.font.color.rgb=G
    s.notes_slide.notes_text_frame.text=(talk.replace('"','') + ('\n\nIZVORI:\n'+src if src else ''))
out=sys.argv[1]; prs.save(out); print(len(prs.slides._sldIdLst),"slides")
