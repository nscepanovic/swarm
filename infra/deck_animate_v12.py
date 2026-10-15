"""Dodaje animacije na klik (fade in) u deck v12, da slajdovi prate govor iz
content/colosseum/pitch-video.md. Ne menja tekst ni raspored.

Pokretanje: python3 infra/deck_animate_v12.py <ulaz.pptx> <izlaz.pptx>

STEPS: po slajdu lista klikova; svaki klik je lista indeksa oblika (redosled
kojim ih pravi infra/deck_build_v12.py) koji se pojavljuju zajedno. Oblici
koji nisu ni u jednom kliku vide se od pocetka (naslov, izvor). Ako se u build
skripti doda ili izbaci oblik, indeksi ovde moraju da se poprave."""
import sys
from lxml import etree
from pptx import Presentation

def r(a, b): return list(range(a, b + 1))

STEPS = {
    # 1 naslovna: bez animacije
    2: [r(1, 6),                    # 150 people, 300 smart hives, 27 hours
        r(7, 9),                    # On Solana: the raise and the token
        r(10, 14)],                 # farma u Bosni, slike
    3: [[2, 3, 4, 5, 10, 13],       # beekeeper -> packer, $2.45
        [6, 7, 8, 9, 11, 12, 14, 15],  # store, customer, $7.15
        [16]],                      # removes the middleman
    4: [r(1, 3),                    # RENT
        r(4, 6),                    # BUY
        [7, 8],                     # who sells
        [9, 10],                    # how we earn
        [11, 12]],                  # trust
    5: [r(1, 4),                    # 1 community funds a farm
        [21] + r(5, 8),             # 2 farm produces
        [22] + r(9, 12),            # 3 marketplace
        [23] + r(13, 16),           # 4 sales, next farm
        [24, 26] + r(17, 20) + [25],  # 5 new farm, krug se zatvara
        [27],                       # community is our first buyer
        [28]],                      # step 1 is done
    6: [r(1, 5),                    # Texas
        r(6, 9),                    # EU
        r(10, 12),                  # vision 1
        r(13, 15),                  # vision 2
        r(16, 18)],                 # vision 3
    7: [r(6, 19),                   # adopt a hive or a tree
        r(20, 27),                  # trade between businesses
        r(28, 41),                  # sell farmland
        r(42, 55),                  # HiveBits red i legenda
        [56]],                      # one loop
    8: [r(1, 4),                    # team
        r(7, 9),                    # Colosseum before
        [5, 6]],                    # next
}

P = 'http://schemas.openxmlformats.org/presentationml/2006/main'

def effect(ids, spid, first):
    node = 'clickEffect' if first else 'withEffect'
    return (
        f'<p:par><p:cTn id="{next(ids)}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" nodeType="{node}">'
        f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
        f'<p:set><p:cBhvr><p:cTn id="{next(ids)}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
        f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
        f'<p:to><p:strVal val="visible"/></p:to></p:set>'
        f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{next(ids)}" dur="400"/>'
        f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
        f'</p:childTnLst></p:cTn></p:par>')

def timing(clicks):
    ids = iter(range(3, 100000))
    body = ''
    for spids in clicks:
        effs = ''.join(effect(ids, sp, i == 0) for i, sp in enumerate(spids))
        body += (
            f'<p:par><p:cTn id="{next(ids)}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{next(ids)}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'{effs}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
    return (
        f'<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
        f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{body}</p:childTnLst></p:cTn>'
        f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
        f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
        f'</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')

prs = Presentation(sys.argv[1])
total = 0
for n, slide in enumerate(prs.slides, 1):
    if n not in STEPS: continue
    shapes = list(slide.shapes)
    used = [i for c in STEPS[n] for i in c]
    assert len(used) == len(set(used)) and max(used) < len(shapes), f'slajd {n}: los indeks'
    clicks = [[shapes[i].shape_id for i in c] for c in STEPS[n]]
    sld = slide._element
    for old in sld.findall(f'{{{P}}}timing'): sld.remove(old)
    sld.append(etree.fromstring(timing(clicks)))
    total += len(clicks)
    print(f'slajd {n}: {len(clicks)} klikova')
prs.save(sys.argv[2]); print('ukupno klikova:', total)
