"""Colosseum deck v5 (otvara rejzom, tri koraka marketplace-a): tamna pozadina, zuti naslovi, po Nemanjinom Google Slides decku
("Colosseum Deck V1 - Marketplace"). Pokretanje: python3 infra/deck_build_v5.py <izlaz.pptx>
Ikonice: Noto Emoji v2.042 (Google, Apache 2.0), u infra/deck_icons/."""
import sys, math, os
ICONS=os.path.join(os.path.dirname(os.path.abspath(__file__)),'deck_icons')
def icon_pic(s,name,cx,y,size=1.25):
    s.shapes.add_picture(os.path.join(ICONS,name+'.png'),Inches(cx-size/2),Inches(y),Inches(size),Inches(size))
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

Y=RGBColor(0xEF,0xB7,0x3E); W=RGBColor(0xFF,0xFF,0xFF); G=RGBColor(0x9A,0x9A,0xA0)
R=RGBColor(0xD9,0x60,0x5A); L=RGBColor(0x8A,0x8A,0x90)
FONT='Poppins'
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
C,Lf,Rt=PP_ALIGN.CENTER,PP_ALIGN.LEFT,PP_ALIGN.RIGHT

def slide(notes=''):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    f=s.background.fill; f.gradient(); f.gradient_angle=315
    f.gradient_stops[0].color.rgb=RGBColor(0x1C,0x1C,0x1F); f.gradient_stops[1].color.rgb=RGBColor(0x38,0x37,0x3C)
    if notes: s.notes_slide.notes_text_frame.text=notes
    return s

def tb(s,x,y,w,h,runs,size=20,color=W,bold=False,align=C,anchor=MSO_ANCHOR.TOP,font=FONT):
    """runs: string, ili lista (tekst, boja) za vise boja u jednom redu; '\n' pravi novi pasus."""
    t=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)).text_frame
    t.word_wrap=True; t.vertical_anchor=anchor
    if isinstance(runs,str): runs=[(runs,color)]
    p=t.paragraphs[0]; p.alignment=align
    for text,col in runs:
        parts=text.split('\n')
        for i,part in enumerate(parts):
            if i>0: p=t.add_paragraph(); p.alignment=align
            if not part: continue
            r=p.add_run(); r.text=part; r.font.size=Pt(size); r.font.bold=bold
            r.font.color.rgb=col
            if font: r.font.name=font
    return t

def title(s,text): tb(s,0.8,0.55,11.7,0.9,text,32,Y,True,Lf)
def footer(s,text): tb(s,0.8,6.85,11.7,0.4,text,10,G,False,Lf)

def arrow(s,x1,y,x2):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y),Inches(x2),Inches(y))
    c.line.color.rgb=L; c.line.width=Pt(1.25)
    ln=c.line._get_or_add_ln(); ln.append(etree.SubElement(ln,qn('a:tailEnd'),type='triangle'))

NODES=[(1.7,'farmer','Farmer'),(5.0,'factory','Food company'),(8.3,'store','Store'),(11.6,'customer','Customer')]
def chain(s,y=1.5):
    for x,icon,label in NODES:
        icon_pic(s,icon,x,y)
        tb(s,x-1.4,y+1.35,2.8,0.5,label,18,W,True)
    for i in range(3):
        arrow(s,NODES[i][0]+0.9,y+0.7,NODES[i+1][0]-0.9)


def xs(s,y=1.5):
    for i in range(3):
        mid=(NODES[i][0]+NODES[i+1][0])/2
        tb(s,mid-0.4,y+0.33,0.8,0.7,'X',30,R,True)

def cols3(s,items,y_icon=1.9,y_head=3.4,y_sub=4.9,head_size=20):
    for i,(icon,head,sub) in enumerate(items):
        x=0.8+i*4.05
        icon_pic(s,icon,x+1.8,y_icon)
        tb(s,x,y_head,3.6,1.3,head,head_size,W,True)
        tb(s,x,y_sub,3.6,1.2,sub,16,G)

# 1 dokaz
s=slide('We are HiveBits. We asked for $140K to build a farm. In 27 hours, 150 people offered $243K.\n\n'
        'We build the first community owned farm.')
tb(s,0.8,1.3,11.7,1.9,'$243K',110,Y,True,Lf)
tb(s,0.8,3.5,11.7,0.8,'offered by 150 people in 27 hours.',34,W,True,Lf)
tb(s,0.8,4.4,11.7,0.6,'We asked for $140K.',24,W,False,Lf)
tb(s,0.8,6.0,7.0,0.5,'HiveBits. The first community owned farm.',20,Y,True,Lf)
footer(s,'Source: HiveBits on the MetaDAO launchpad, $NECTAR. Raise goal $140K, 173%, 150 funders.')

# 2 kontekst
s=slide('Have you heard of the term "safe food"? It is food where you can check where it comes from and what is inside.\n\n'
        'Most people know only two kinds. Conventional food and organic food.\n\nDo you know what you eat?')
tb(s,0.8,2.0,3.6,1.4,'Conventional\nfood',28,W,True)
tb(s,4.87,2.0,3.6,1.4,'Safe food',28,Y,True)
tb(s,8.93,2.0,3.6,1.4,'Organic food',28,W,True)
tb(s,2.5,4.6,8.33,1.0,'Do you know what you eat?',30,Y,True)

# 3 problem
s=slide('Take milk in Bosnia. The farmer gets only 30 percent of the final price.\n\n'
        'When cheap milk comes from import, the dairy pays him less.\n\n'
        'This year his price went down 25 percent.\n\n'
        'The result is that the customer pays three times more than the farmer gets, for a processed product.')
chain(s)
for i,price in enumerate(['$0.34 per liter','$0.73 per liter','$1.09 per liter']):
    mid=(NODES[i][0]+NODES[i+1][0])/2
    tb(s,mid-1.2,1.65,2.4,0.4,price,14,G)
tb(s,1,4.6,11.33,0.9,[("Farmer's price: ",W),('25%',R),(' less than last year.',W)],28,W,True)
footer(s,'Milk in Bosnia, 2026. Farmer price: Viteski.ba, 29 Sep 2026. Shelf price: 072info, 29 Jul 2026. Middle price is our estimate.')

# 4 resenje
s=slide("People know that food from the farmer is better. But they don't know how to find a farmer they can trust. So they go to the store.\n\n"
        'HiveBits solves that. You buy food directly from farmers you can trust.\n\n'
        'For the same liter of milk, the farmer gets about 51 cents instead of 34. You pay about 75 instead of $1.09.')
title(s,'Buy food from trusted farmers.')
Y4=1.9
chain(s,Y4); xs(s,Y4)
x1,x2,y0,depth=NODES[0][0],NODES[3][0],Y4+2.45,1.4
pts=[]
for k in range(0,61):
    a=math.pi*k/60
    pts.append((Inches((x1+x2)/2-(x2-x1)/2*math.cos(a)),Inches(y0+depth*math.sin(a))))
fb=s.shapes.build_freeform(pts[0][0],pts[0][1]); fb.add_line_segments(pts[1:],close=False)
arc=fb.convert_to_shape(); arc.fill.background(); arc.line.color.rgb=Y; arc.line.width=Pt(1.5)
ln=arc.line._get_or_add_ln(); etree.SubElement(ln,qn('a:tailEnd'),type='triangle')
tb(s,4.67,Y4+2.8,4.0,0.5,'HiveBits',24,Y,True)
tb(s,4.17,Y4+3.3,5.0,0.4,'about $0.75 per liter',16,W)
tb(s,0.3,Y4+1.85,2.8,0.4,'$0.51, not $0.34',15,Y,True)
tb(s,10.2,Y4+1.85,2.8,0.4,'$0.75, not $1.09',15,Y,True)
footer(s,'Direct price is our estimate: delivery twice a week, 400 liters per trip.')

# 5 tri koraka
s=slide('We already did the last step. People invested in our farm.\n\n'
        'The marketplace does the same for other farmers. We find the farm, and people can invest through us.\n\n'
        'Before that you can rent a hive or buy ahead. And you can simply buy the product, like in any shop.')
title(s,'One marketplace, three steps')
cols3(s,[('customer','1. Buy.','Food from farmers we know.'),
         ('bee','2. Rent or buy ahead.','A hive, a cow, the next harvest.'),
         ('seedling','3. Invest in the farm.','150 people, $243K offered.')])
tb(s,8.9,5.6,3.6,0.5,'We did this first.',18,Y,True)

# 6 poverenje
s=slide('Why can you trust the farm. The community owns it. Every hive has our sensor. We track the weight, the temperature and the humidity.\n\n'
        'Your money is locked on Solana, in stablecoins. The farmer gets it when he delivers the food.\n\n'
        'For other farms we start with farmers we know.')
title(s,'Why you can trust the farm')
cols3(s,[('handshake','The community owns the farm.','150 people funded it.'),
         ('antenna','Every hive has our sensor.','Weight, temperature, humidity.'),
         ('lock','The money is locked until delivery.','On Solana, in stablecoins.')])
tb(s,0.8,6.0,11.7,0.5,'We start with farmers we know.',18,W,False,Lf)

# 7 go to market
s=slide('We start with what we know. Our own farm with 300 smart hives is funded.\n\n'
        'Next is the second farm, with a beekeeper we know. It is funded the same way we funded ours. Why should people stay without good honey.\n\n'
        'When this works for honey, we bring other farmers.')
title(s,'Go to market')
gtm=[('1','Our own farm.','300 smart hives, funded in 27 hours.'),('2','The second farm.','Funded through us, with a beekeeper we know.'),('3','Other farmers.','Same platform, same proof.')]
for i,(n,head,sub) in enumerate(gtm):
    x=0.8+i*4.05
    tb(s,x,1.8,3.6,1.5,n,80,Y,True,Lf)
    tb(s,x,3.6,3.7,1.1,head,22,W,True,Lf)
    tb(s,x,4.8,3.7,1.0,sub,18,G,False,Lf)

# 8 konkurencija
s=slide('Many teams tried a farm marketplace here. None of them won. They were all software with no farmer behind it.\n\n'
        'We are farmers. We have our hives and our sensors, and 150 people already backed us.')
title(s,'Competition')
tb(s,0.8,1.7,5.6,0.5,'Others',20,G,True,Lf); tb(s,6.9,1.7,5.6,0.5,'HiveBits',20,Y,True,Lf)
rows=[('Farm marketplaces entered Colosseum many times. None won a prize.','We have our own farm, our own hardware and 150 backers.'),
      ('CrowdFarming takes 16 to 32 percent.','We take 5 percent.'),
      ('3Bee, Pollenity, Honeyverse sell only their own hives, off chain.','Any farmer can sell. The payment is on Solana.')]
for i,(a,b) in enumerate(rows):
    y=2.4+i*1.35
    ln_=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(0.8),Inches(y-0.1),Inches(12.5),Inches(y-0.1)); ln_.line.color.rgb=RGBColor(0x55,0x55,0x5A); ln_.line.width=Pt(0.75)
    tb(s,0.8,y,5.6,1.1,a,18,W,False,Lf); tb(s,6.9,y,5.6,1.1,b,18,W,True,Lf)
footer(s,'Source: Colosseum Copilot, four hackathons from March 2024 to September 2025, checked 10 Oct 2026. CrowdFarming economic transparency page. 5%: HiveBits.')

# 9 zarada i trziste
s=slide('We earn in two places. A fee when a farm is funded, and 5 percent of every sale after that.\n\n'
        'We could take only $140K. The other $103K had no farm to go to. That is why we build the marketplace.\n\n'
        'Later the same farmers sell to shops and restaurants. AgriDex already does farm trade on Solana, in bulk.')
title(s,'How we earn')
tb(s,0.8,1.8,5.6,0.5,'Today',20,Y,True,Lf)
tb(s,0.8,2.3,5.6,0.6,'A fee when a farm is funded.',22,W,True,Lf)
tb(s,0.8,2.95,5.6,0.6,'5% of every sale.',22,W,True,Lf)
tb(s,0.8,4.0,5.6,0.5,'Next',20,Y,True,Lf)
tb(s,0.8,4.5,5.6,1.0,'The same farmers sell to shops and restaurants.',22,W,True,Lf)
tb(s,0.8,5.7,5.6,0.8,'AgriDex does this on Solana in bulk. $9M+ in trades.',16,G,False,Lf)
tb(s,7.2,1.6,5.4,1.2,'$103K',64,Y,True,Lf); tb(s,7.2,2.9,5.4,0.9,'offered and not taken. Demand is bigger than one farm.',18,W,False,Lf)
tb(s,7.2,4.2,5.4,1.2,'$839M',64,Y,True,Lf); tb(s,7.2,5.5,5.4,0.6,'US households spend on honey every year',18,W,False,Lf)
footer(s,'Source: MetaDAO launchpad ($243K offered, $140K goal). NielsenIQ for National Honey Board, 52 weeks to 2 Nov 2024. The Defiant on AgriDex.')

# 10 tim i ask
s=slide('We are a team of beekeepers and engineers.\n\n'
        'We are here to win and to join the Colosseum accelerator. Next we bring more farmers and we build better ways to check them.')
title(s,'Team')
team=[('Nemanja Šćepanović','CEO, beekeeper','10 years in engineering and IoT.'),
      ('Siniša Kežić','COO, chief beekeeper','Runs the farm build.')]
for i,(n,role,bio) in enumerate(team):
    x=0.8+i*4.05
    tb(s,x,1.5,3.7,1.0,n,22,W,True,Lf,MSO_ANCHOR.BOTTOM); tb(s,x,2.6,3.7,0.5,role,16,Y,True,Lf); tb(s,x,3.1,3.7,0.6,bio,16,G,False,Lf)
tb(s,0.8,3.9,11.7,0.5,'Raised on MetaDAO. Backed by Superteam Balkan.',16,G,False,Lf)
tb(s,0.8,4.7,11.7,0.8,'We are here to win and join the Colosseum accelerator.',22,Y,True,Lf)
tb(s,0.8,5.5,11.7,0.5,'Next: more farmers, and better ways to check them.',18,W,False,Lf)
tb(s,0.8,6.3,11.7,0.5,'hivebits.io',20,W,True,Lf)

prs.save(sys.argv[1]); print(len(prs.slides._sldIdLst),'slides')
